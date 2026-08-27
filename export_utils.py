import csv
from typing import List, Tuple, Any
from models import Component
from database import Database

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False


def _calc_display_width(val: Any) -> int:
    """Calculate the precise display width needed for text to prevent truncation in Excel."""
    if val is None:
        return 0
    text = str(val)
    width = 0.0
    for ch in text:
        # Non-ASCII and Persian characters occupy more visual width in Excel
        if ord(ch) > 127:
            width += 1.45
        else:
            width += 1.05
    return int(width)


def export_to_excel(components: List[Component], file_path: str) -> str:
    """Generate engineering Excel workbook with auto-fitted dimensions and styling."""
    if not EXCEL_AVAILABLE:
        print("[WARN] [EXPORT] openpyxl not found. Falling back to CSV export.")
        csv_path = file_path.replace('.xlsx', '.csv')
        export_to_csv(components, csv_path)
        return csv_path

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.sheet_view.rightToLeft = True  # Enable RTL view for Persian columns

    # Header and cell styling definitions
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="0F172A")
    mono_font = Font(name="Consolas", size=10, bold=True, color="0F172A")
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    headers = [
        "ردیف",
        "نام قطعه",
        "پارت‌نامبر / مقدار",
        "پکیج / فوت‌پرینت",
        "دسته‌بندی",
        "موجودی کل",
        "حداقل هشدار",
        "کشوها و تفکیک موجودی",
        "وضعیت",
        "توضیحات"
    ]
    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_num, cell in enumerate(ws[1], 1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = border_thin

    # Status column color fills
    fill_ok = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")      # Matte Green
    fill_low = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")     # Matte Amber
    fill_empty = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")   # Matte Rose/Red
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")   # Subtle Zebra Stripe

    for idx, comp in enumerate(components, 1):
        row_data = [
            idx,
            comp.name,
            comp.value,
            comp.package or "—",
            comp.category,
            comp.quantity,
            comp.min_alert,
            comp.formatted_drawers_detailed,
            comp.status_label,
            comp.description or "—"
        ]
        ws.append(row_data)
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 24

        is_even_row = (idx % 2 == 0)

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.border = border_thin
            cell.font = data_font

            # Cell alignment rules
            if col_idx in (1, 4, 5, 6, 7, 8, 9):
                cell.alignment = align_center
            elif col_idx in (2, 10):
                cell.alignment = align_right
            elif col_idx == 3:
                cell.alignment = align_center
                cell.font = mono_font

            # Alternate row zebra coloring
            if is_even_row:
                cell.fill = fill_zebra

            # Status column color coding (Column 9)
            if col_idx == 9:
                if comp.stock_status == "OK":
                    cell.fill = fill_ok
                elif comp.stock_status == "LOW":
                    cell.fill = fill_low
                else:
                    cell.fill = fill_empty

    # Auto-fit column widths
    col_min_widths = {
        1: 8,   # Row index
        2: 22,  # Name
        3: 26,  # Value / Part Number
        4: 16,  # Package
        5: 20,  # Category
        6: 14,  # Total Quantity
        7: 14,  # Min Alert
        8: 26,  # Drawers
        9: 14,  # Status
        10: 32  # Description
    }

    for col_idx, col in enumerate(ws.columns, 1):
        max_len = max(_calc_display_width(cell.value) for cell in col)
        col_letter = get_column_letter(col_idx)
        min_w = col_min_widths.get(col_idx, 14)
        ws.column_dimensions[col_letter].width = max(max_len + 5, min_w)

    wb.save(file_path)
    print(f"[INFO] [EXPORT] Professional Excel export saved to: {file_path}")
    return file_path


def export_to_csv(components: List[Component], file_path: str) -> str:
    """Export components to a standard UTF-8 BOM encoded CSV file."""
    with open(file_path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow([
            "ID", "Name", "Part_Number", "Package", "Category", 
            "Total_Quantity", "Min_Alert", "Drawer_Stocks", "Status", "Description"
        ])
        for c in components:
            writer.writerow([
                c.id,
                c.name,
                c.value,
                c.package,
                c.category,
                c.quantity,
                c.min_alert,
                c.drawers,
                c.status_label,
                c.description
            ])
    print(f"[INFO] [EXPORT] CSV export saved to: {file_path}")
    return file_path


def import_from_csv(file_path: str, db: Database) -> Tuple[int, int]:
    """Bulk import components from a CSV file into the inventory database."""
    success_count = 0
    fail_count = 0
    with open(file_path, 'r', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        for row in reader:
            try:
                name = row.get('Name') or row.get('نام قطعه') or 'قطعه جدید'
                val = row.get('Part_Number') or row.get('Value') or row.get('مقدار') or ''
                pkg = row.get('Package') or row.get('پکیج') or ''
                cat = row.get('Category') or row.get('دسته‌بندی') or 'Other'
                qty = int(row.get('Total_Quantity') or row.get('Quantity') or row.get('موجودی') or 0)
                min_a = int(row.get('Min_Alert') or row.get('MinAlert') or row.get('حداقل موجودی') or 10)
                drawers = row.get('Drawer_Stocks') or row.get('Drawers') or row.get('کشوها') or ''
                desc = row.get('Description') or row.get('توضیحات') or ''

                comp = Component(
                    name=name,
                    value=val,
                    package=pkg,
                    category=cat,
                    quantity=qty,
                    min_alert=min_a,
                    drawers=drawers,
                    description=desc
                )
                db.add_component(comp)
                success_count += 1
            except Exception as e:
                print(f"[WARN] [IMPORT] Failed to parse row: {e}")
                fail_count += 1
    print(f"[INFO] [IMPORT] CSV import finished: {success_count} succeeded, {fail_count} failed")
    return success_count, fail_count
