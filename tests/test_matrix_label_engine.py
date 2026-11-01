import os
import pytest
from bom_auditor import BOMItem
from ui.printing.matrix_label_engine import MatrixLabelEngine


def test_matrix_label_engine_excel_generation(tmp_path):
    items = [
        BOMItem(1, comment="100uF", description="Capacitor D Case 100uF 16VDC", designator="C1", footprint="SMDCAPCASED", lib_ref="SMD Tantalum Capacitor", unit_qty=1),
        BOMItem(2, comment="4.7uF", description="Capacitor 0805 4.7uF 50VDC", designator="C2, C23", footprint="SMDCAP0805", lib_ref="SMD Ceramic Capacitor", unit_qty=2),
    ]

    out_file = str(tmp_path / "FC_IO_V1.3_Matrix.xlsx")
    saved = MatrixLabelEngine.export_excel(items, "FC_IO_V1.3", out_file)
    assert os.path.exists(saved)

    import openpyxl
    wb = openpyxl.load_workbook(saved)
    ws = wb.active
    assert ws.title == "Matrix Labels"
    assert ws.cell(1, 1).value == "FC_IO_V1.3"
    assert ws.cell(1, 2).value == "FC_IO_V1.3_N1"
    assert ws.cell(1, 3).value == "FC_IO_V1.3_N2"

    assert ws.cell(2, 1).value == "Comment"
    assert ws.cell(2, 2).value == "100uF"
    assert ws.cell(2, 3).value == "4.7uF"

    assert ws.cell(7, 1).value == "Quantity"
    assert ws.cell(7, 2).value == "1"
    assert ws.cell(7, 3).value == "2"


def test_matrix_label_engine_html_cut_strip_generation():
    items = [
        BOMItem(1, comment="100uF", description="Capacitor D Case", designator="C1", footprint="SMDCAPCASED", lib_ref="Tantalum", unit_qty=1),
    ]
    html_doc = MatrixLabelEngine.generate_cut_strip_html(items, "FC_IO_V1.3")
    assert "FC_IO_V1.3_N1" in html_doc
    assert "✂ - - - - - - - - - - - - ✂" in html_doc
    assert "100uF" in html_doc
    assert "border: 2px dashed #000000;" in html_doc
