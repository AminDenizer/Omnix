"""
Tests for Dual-Path Mode A Validation: Stock Shortages + Altium BOM.

Governing Requirements:
- Section 3: Path A — Stock Shortages Integration
- Section 4: Path B — Altium BOM Procurement Integration (based on FC_IO_V_1_3(2) / FC_IO_V_1_2 structure)
- Section 5: Preserved Multi-Vendor Contracts Across Both Paths
- Section 6: Separately Identifiable Entry-Point Integration Tests
"""

import os
from typing import Any, Dict, List, Optional
from unittest.mock import MagicMock, patch

import openpyxl
import pytest
from PyQt6.QtWidgets import QApplication, QDialog

from bom_auditor import BOMAuditor, BOMItem, BOMAuditResult
from core.models import AdapterConfig, AdapterStatus, ProductDetails, ProductReference
from core.procurement.bridge import ProcurementBridge
from core.procurement.models import ProcurementMode, ProcurementRequest
from core.procurement.registry import VendorRegistry
from core.procurement.vendor_links import parse_vendor_links
from core.vendor_interface import IVendorAdapter
from excel_loader import ExcelDataLoader
from ui.dialogs.order_preview_dialog import OrderPreviewDialog, VendorSelectionDialog


@pytest.fixture(scope="module")
def qapp():
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    return app


class MockLionAdapter(IVendorAdapter):
    VENDOR_ID = "vendor_lion"
    BASE_URL = "https://lionelectronic.ir"

    def __init__(self, config: Optional[AdapterConfig] = None):
        self._config = config
        self.dispatched_urls: List[str] = []

    @property
    def vendor_id(self) -> str:
        return self.VENDOR_ID

    @property
    def base_url(self) -> str:
        return self.BASE_URL

    def initialize(self, config: Optional[AdapterConfig] = None) -> AdapterStatus:
        return AdapterStatus(vendor_id=self.VENDOR_ID, is_ready=True, is_authenticated=True)

    def search_part(self, query: str) -> Any:
        return None

    def resolve_product(self, url_or_sku: str) -> ProductReference:
        return ProductReference(vendor_id=self.VENDOR_ID, vendor_sku="LION-SKU", pdp_url=url_or_sku)

    def inspect_product(self, product_ref: ProductReference) -> ProductDetails:
        return ProductDetails(
            vendor_id=self.VENDOR_ID,
            vendor_sku=product_ref.vendor_sku,
            pdp_url=product_ref.pdp_url,
            stock_quantity=100,
            can_order=True,
            price_irr=5000,
            mpn=product_ref.vendor_sku,
        )

    def extract_product_data(self, dom_context: Any) -> Any:
        return None

    def inspect_stock(self, product_ref: ProductReference) -> Any:
        return None

    def determine_order_constraints(self, product_ref: ProductReference) -> Any:
        return None

    def add_to_cart(self, product_ref: ProductReference, quantity: int) -> Any:
        self.dispatched_urls.append(product_ref.pdp_url)
        return {"success": True, "staged_quantity": quantity}

    def inspect_cart(self) -> Any:
        return None

    def authenticate(self, credentials: Any = None) -> Any:
        return None

    def validate_session(self) -> Any:
        return None

    def recover_from_error(self, error_context: Any) -> Any:
        return None

    def close(self) -> None:
        pass


class MockJavanAdapter(IVendorAdapter):
    VENDOR_ID = "vendor_javan"
    BASE_URL = "https://javanelec.com"

    def __init__(self, config: Optional[AdapterConfig] = None):
        self._config = config
        self.dispatched_urls: List[str] = []

    @property
    def vendor_id(self) -> str:
        return self.VENDOR_ID

    @property
    def base_url(self) -> str:
        return self.BASE_URL

    def initialize(self, config: Optional[AdapterConfig] = None) -> AdapterStatus:
        return AdapterStatus(vendor_id=self.VENDOR_ID, is_ready=True, is_authenticated=True)

    def search_part(self, query: str) -> Any:
        return None

    def resolve_product(self, url_or_sku: str) -> ProductReference:
        return ProductReference(vendor_id=self.VENDOR_ID, vendor_sku="JAVAN-SKU", pdp_url=url_or_sku)

    def inspect_product(self, product_ref: ProductReference) -> ProductDetails:
        return ProductDetails(
            vendor_id=self.VENDOR_ID,
            vendor_sku=product_ref.vendor_sku,
            pdp_url=product_ref.pdp_url,
            stock_quantity=200,
            can_order=True,
            price_irr=4500,
            mpn=product_ref.vendor_sku,
        )

    def extract_product_data(self, dom_context: Any) -> Any:
        return None

    def inspect_stock(self, product_ref: ProductReference) -> Any:
        return None

    def determine_order_constraints(self, product_ref: ProductReference) -> Any:
        return None

    def add_to_cart(self, product_ref: ProductReference, quantity: int) -> Any:
        self.dispatched_urls.append(product_ref.pdp_url)
        return {"success": True, "staged_quantity": quantity}

    def inspect_cart(self) -> Any:
        return None

    def authenticate(self, credentials: Any = None) -> Any:
        return None

    def validate_session(self) -> Any:
        return None

    def recover_from_error(self, error_context: Any) -> Any:
        return None

    def close(self) -> None:
        pass


@pytest.fixture
def dual_path_registry():
    """Sets up default registry with Mock Lion and Mock Javan adapters with clean teardown."""
    reg = VendorRegistry.get_default_registry()
    reg.clear()
    reg.register_adapter(
        vendor_id="vendor_lion",
        adapter_class=MockLionAdapter,
        domains=["lionelectronic.ir"],
        display_name="Lion Electronic",
        enabled=True,
        activate=True,
    )
    reg.register_adapter(
        vendor_id="vendor_javan",
        adapter_class=MockJavanAdapter,
        domains=["javanelec.com"],
        display_name="Javan Electronic",
        enabled=True,
        activate=True,
    )
    yield reg
    reg.clear()


# =============================================================================
# PATH A: STOCK SHORTAGES INTEGRATION TESTS
# =============================================================================

class TestPathAStockShortagesIntegration:
    """
    Validates Entry Point A: Stock Shortages / Replenishment.
    Verifies that shortage items loaded from Inventory.xlsx flow correctly
    through multi-vendor link resolution, OrderPreviewDialog, user selection,
    and atomic adapter dispatch.
    """

    @pytest.fixture
    def inventory_file_path_a(self, tmp_path):
        """Creates a realistic Inventory workbook for Stock Shortages."""
        inv_path = str(tmp_path / "Inventory_PathA.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventory"
        ws.append([
            "ECODE", "PART NAME", "DISCRIPTION", "Quantity in Stock",
            "FOOTPRINT", "PART NUMBER", "LINK", "Min Stock", "Target Stock"
        ])
        # 1. Single Legacy Lion URL (in shortage)
        ws.append([
            None, "NE555P", "Timer IC DIP-8", 2,
            "DIP-8", "NE555P",
            "https://lionelectronic.ir/products/2070-NE555P", 10, 25
        ])
        # 2. Single Labeled Javan URL (in shortage)
        ws.append([
            None, "T2N7002BK", "N-Channel MOSFET SOT-23", 5,
            "SOT-23", "T2N7002BK",
            "JAVAN=https://www.javanelec.com/shop/product/55407/toshiba-corporation/t2n7002bk", 20, 50
        ])
        # 3. Multi-Vendor Link: Lion and Javan (in shortage)
        ws.append([
            None, "NUP2105LT1G", "Dual Line CAN Bus Protector", 0,
            "SOT-23", "NUP2105LT1G",
            "LION=https://lionelectronic.ir/products/3775-NUP2105LT1G || JAVAN=https://www.javanelec.com/shop/product/55407/toshiba-corporation/t2n7002bk",
            15, 30
        ])
        # 4. In Stock item (Not a shortage, should be excluded from shortages)
        ws.append([
            None, "LM358N", "Dual Op-Amp", 100,
            "DIP-8", "LM358N",
            "https://lionelectronic.ir/products/1932-LM358N", 10, 20
        ])
        # 5. Invalid / Unsupported URL (in shortage)
        ws.append([
            None, "UNKNOWN_CHIP", "Special ASIC", 0,
            "QFN-32", "UNKNOWN_CHIP",
            "https://unsupported-supplier.com/item/999", 5, 10
        ])
        # 6. Blank Link (in shortage)
        ws.append([
            None, "RES_10K", "10k Resistor", 0,
            "0805", "RC0805-10K",
            "", 50, 100
        ])
        wb.save(inv_path)
        wb.close()
        return inv_path

    def test_path_a_loader_classifies_shortages_accurately(self, inventory_file_path_a, dual_path_registry):
        loader = ExcelDataLoader()
        ok, msg = loader.load_file(inventory_file_path_a)
        assert ok is True

        shortages = loader.get_shortage_items(registry=dual_path_registry)
        # Total rows in shortage: 5 (items 1, 2, 3, 5, 6; item 4 has stock=100 > min=10)
        assert len(shortages) == 5

        # Check Multi-Vendor shortage
        nup = next(s for s in shortages if s["part_number"] == "NUP2105LT1G")
        assert nup["is_multi_vendor"] is True
        assert nup["link_type"] == "MULTI"
        assert nup["is_orderable"] is True
        assert len(nup["vendor_links"]) == 2
        assert "vendor_lion" in nup["vendor_links"]
        assert "vendor_javan" in nup["vendor_links"]

        # Check Single Lion legacy shortage
        ne555 = next(s for s in shortages if s["part_number"] == "NE555P")
        assert ne555["is_multi_vendor"] is False
        assert ne555["link_type"] == "LION"
        assert ne555["vendor_id"] == "vendor_lion"
        assert ne555["is_orderable"] is True

        # Check Single Javan labeled shortage
        mos = next(s for s in shortages if s["part_number"] == "T2N7002BK")
        assert mos["is_multi_vendor"] is False
        assert mos["link_type"] == "JAVAN"
        assert mos["vendor_id"] == "vendor_javan"
        assert mos["is_orderable"] is True

        # Check Unsupported URL shortage
        unsupported = next(s for s in shortages if s["part_number"] == "UNKNOWN_CHIP")
        assert unsupported["is_orderable"] is False
        assert unsupported["link_type"] == "INVALID"

        # Check Blank Link shortage
        blank = next(s for s in shortages if s["part_number"] == "RC0805-10K")
        assert blank["is_orderable"] is False
        assert blank["link_type"] == "NONE"

    def test_path_a_single_link_bypasses_selection_dialog(self, qapp, inventory_file_path_a, dual_path_registry):
        """Single-link items must bypass VendorSelectionDialog completely."""
        loader = ExcelDataLoader()
        loader.load_file(inventory_file_path_a)
        shortages = loader.get_shortage_items(registry=dual_path_registry)
        single_link_items = [s for s in shortages if s["part_number"] in ("NE555P", "T2N7002BK")]

        dlg = OrderPreviewDialog(all_shortages=single_link_items, registry=dual_path_registry)
        assert len(dlg.orderable_items) == 2
        assert len(dlg.multi_vendor_items) == 0

        with patch.object(VendorSelectionDialog, "exec") as mock_dlg_exec:
            with patch("ui.dialogs.order_preview_dialog.OrderWorkerThread") as mock_worker_cls:
                with patch("ui.dialogs.order_preview_dialog.settings.get", return_value="dummy_val"):
                    dlg._execute_mode_a_order()
                    # Dialog must NEVER be instantiated or executed for single-link items
                    mock_dlg_exec.assert_not_called()
                    assert mock_worker_cls.called

    def test_path_a_multi_link_prompts_and_dispatches_lion(self, qapp, inventory_file_path_a, dual_path_registry):
        """User selecting Lion from multi-vendor prompt dispatches only Lion URL."""
        loader = ExcelDataLoader()
        loader.load_file(inventory_file_path_a)
        shortages = loader.get_shortage_items(registry=dual_path_registry)
        multi_item = [s for s in shortages if s["part_number"] == "NUP2105LT1G"]

        dlg = OrderPreviewDialog(all_shortages=multi_item, registry=dual_path_registry)
        assert len(dlg.multi_vendor_items) == 1

        def select_lion(*args, **kwargs):
            dlg.multi_vendor_items[0]["selected_vendor_id"] = "vendor_lion"
            dlg.multi_vendor_items[0]["selected_url"] = dlg.multi_vendor_items[0]["vendor_links"]["vendor_lion"]
            return QDialog.DialogCode.Accepted

        with patch.object(VendorSelectionDialog, "exec", side_effect=select_lion) as mock_exec:
            with patch("ui.dialogs.order_preview_dialog.OrderWorkerThread") as mock_worker_cls:
                with patch("ui.dialogs.order_preview_dialog.settings.get", return_value="mock_user"):
                    dlg._execute_mode_a_order()
                    assert mock_exec.called
                    assert mock_worker_cls.called

                    # Verify dispatched item
                    dispatched_items = mock_worker_cls.call_args[1]["items"]
                    req = ProcurementBridge.item_to_request(dispatched_items[0], registry=dual_path_registry)
                    assert req.explicit_vendor_id == "vendor_lion"
                    assert req.explicit_vendor_url == "https://lionelectronic.ir/products/3775-NUP2105LT1G"
                    assert "javanelec" not in req.explicit_vendor_url
                    assert "||" not in req.explicit_vendor_url

    def test_path_a_multi_link_prompts_and_dispatches_javan(self, qapp, inventory_file_path_a, dual_path_registry):
        """User selecting Javan from multi-vendor prompt dispatches only Javan URL."""
        loader = ExcelDataLoader()
        loader.load_file(inventory_file_path_a)
        shortages = loader.get_shortage_items(registry=dual_path_registry)
        multi_item = [s for s in shortages if s["part_number"] == "NUP2105LT1G"]

        dlg = OrderPreviewDialog(all_shortages=multi_item, registry=dual_path_registry)

        def select_javan(*args, **kwargs):
            dlg.multi_vendor_items[0]["selected_vendor_id"] = "vendor_javan"
            dlg.multi_vendor_items[0]["selected_url"] = dlg.multi_vendor_items[0]["vendor_links"]["vendor_javan"]
            return QDialog.DialogCode.Accepted

        with patch.object(VendorSelectionDialog, "exec", side_effect=select_javan) as mock_exec:
            with patch("ui.dialogs.order_preview_dialog.OrderWorkerThread") as mock_worker_cls:
                dlg._execute_mode_a_order()
                assert mock_exec.called
                assert mock_worker_cls.called

                dispatched_items = mock_worker_cls.call_args[1]["items"]
                req = ProcurementBridge.item_to_request(dispatched_items[0], registry=dual_path_registry)
                assert req.explicit_vendor_id == "vendor_javan"
                assert req.explicit_vendor_url == "https://www.javanelec.com/shop/product/55407/toshiba-corporation/t2n7002bk"
                assert "lionelectronic" not in req.explicit_vendor_url

    def test_path_a_cancellation_prevents_dispatch(self, qapp, inventory_file_path_a, dual_path_registry):
        """Cancelling the vendor selection dialog must abort procurement without dispatching."""
        loader = ExcelDataLoader()
        loader.load_file(inventory_file_path_a)
        shortages = loader.get_shortage_items(registry=dual_path_registry)
        multi_item = [s for s in shortages if s["part_number"] == "NUP2105LT1G"]

        dlg = OrderPreviewDialog(all_shortages=multi_item, registry=dual_path_registry)

        with patch.object(VendorSelectionDialog, "exec", return_value=QDialog.DialogCode.Rejected) as mock_exec:
            with patch("ui.dialogs.order_preview_dialog.OrderWorkerThread") as mock_worker_cls:
                dlg._execute_mode_a_order()
                assert mock_exec.called
                # Worker thread must NOT be called on cancellation
                mock_worker_cls.assert_not_called()
                assert dlg.worker is None

    def test_path_a_atomic_adapter_resolution_end_to_end(self, inventory_file_path_a, dual_path_registry):
        """The resolved URL from Path A reaches resolve_adapter_for_url atomically."""
        loader = ExcelDataLoader()
        loader.load_file(inventory_file_path_a)
        shortages = loader.get_shortage_items(registry=dual_path_registry)

        nup = next(s for s in shortages if s["part_number"] == "NUP2105LT1G")
        nup["selected_vendor_id"] = "vendor_lion"
        nup["selected_url"] = nup["vendor_links"]["vendor_lion"]

        req = ProcurementBridge.item_to_request(nup, registry=dual_path_registry)
        adapter, reg_info = dual_path_registry.resolve_adapter_for_url(req.explicit_vendor_url)
        assert adapter is not None
        assert adapter.vendor_id == "vendor_lion"
        assert reg_info.vendor_id == "vendor_lion"
        assert isinstance(adapter, MockLionAdapter)

        # Dispatch via adapter add_to_cart
        p_ref = adapter.resolve_product(req.explicit_vendor_url)
        adapter.add_to_cart(p_ref, req.required_quantity)
        assert req.explicit_vendor_url in adapter.dispatched_urls


# =============================================================================
# PATH B: ALTIUM BOM PROCUREMENT INTEGRATION TESTS
# =============================================================================

class TestPathBBOMProcurementIntegration:
    """
    Validates Entry Point B: Altium BOM Procurement.
    Mirrors FC_IO_V_1_3(2).xlsx / FC_IO_V_1_2.xlsx canonical parameters and verifies
    BOM import, reconciliation with Inventory datastore, multi-vendor link parsing,
    OrderPreviewDialog integration, user selection, and atomic adapter dispatch.
    """

    @pytest.fixture
    def bom_and_inventory_path_b(self, tmp_path):
        """
        Creates a faithful Altium BOM conforming to FC_IO_V_1_2.xlsx / FC_IO_V_1_3(2).xlsx
        and an inventory workbook containing multi-vendor, single-link, and missing items.
        """
        inv_path = str(tmp_path / "Inventory_PathB.xlsx")
        inv_wb = openpyxl.Workbook()
        ws_inv = inv_wb.active
        ws_inv.title = "Inventory"
        ws_inv.append([
            "ECODE", "PART NAME", "DISCRIPTION", "Quantity in Stock",
            "FOOTPRINT", "PART NUMBER", "LINK"
        ])
        # 1. TAJD107K016RNJ: Multi-Vendor (Lion + Javan), in stock=1 (needed 5 -> shortage=4)
        ws_inv.append([
            None, "TAJD107K016RNJ", "Capacitor D Case 100uF 16VDC 10%", 1,
            "SMDCAPCASED", "TAJD107K016RNJ",
            "LION=https://lionelectronic.ir/products/3775-TAJD107K016RNJ || JAVAN=https://www.javanelec.com/shop/product/55407/toshiba-corporation/t2n7002bk"
        ])
        # 2. C2012X5R1H475K125AB: Single Javan Link, in stock=0 (needed 10 -> shortage=10)
        ws_inv.append([
            None, "C2012X5R1H475K125AB", "Capacitor 0805 4.7uF 50VDC 10%", 0,
            "SMDCAP0805", "C2012X5R1H475K125AB",
            "JAVAN=https://www.javanelec.com/shop/product/55407/toshiba-corporation/t2n7002bk"
        ])
        # 3. GRM188R71H104KA93D: Single Lion Legacy Link, in stock=50 (needed 190 -> shortage=140)
        ws_inv.append([
            None, "GRM188R71H104KA93D", "Capacitor 0603 100nF 50VDC 10%", 50,
            "SMDCAP0603", "GRM188R71H104KA93D",
            "https://lionelectronic.ir/products/2070-GRM188R71H104KA93D"
        ])
        # 4. BLANK_STORE_LINK_PART: In DB but link is empty, in stock=0 (needed 5 -> shortage=5)
        ws_inv.append([
            None, "BLANK_LINK_CAP", "Capacitor with no link", 0,
            "0805", "BLANK_LINK_PART",
            ""
        ])
        # 5. UNSUPPORTED_URL_PART: In DB but link domain is not in registry
        ws_inv.append([
            None, "UNSUPPORTED_PART", "Unsupported store part", 0,
            "SOIC-8", "UNSUPPORTED_PART",
            "https://unsupportedstore.com/item/1"
        ])
        inv_wb.save(inv_path)
        inv_wb.close()

        # Create Altium BOM mimicking FC_IO_V_1_2.xlsx
        bom_path = str(tmp_path / "FC_IO_V_1_3(2).xlsx")
        bom_wb = openpyxl.Workbook()
        ws_bom = bom_wb.active
        ws_bom.title = "FC_IO_V_1_3"
        ws_bom.append([
            "Designator", "Comment", "Description", "Designator", "Footprint",
            "LibRef", "Quantity", "Part Number", "Min Stock", "Target Stock"
        ])
        # N1: TAJD107K016RNJ (1 per board)
        ws_bom.append([
            "N1", "100uF", "Capacitor D Case 100uF 16VDC 10%", "C1", "SMDCAPCASED",
            "SMD Tantalum Capacitor", 1, "TAJD107K016RNJ", 0, 0
        ])
        # N2: C2012X5R1H475K125AB (2 per board)
        ws_bom.append([
            "N2", "4.7uF", "Capacitor 0805 4.7uF 50VDC 10%", "C2, C23", "SMDCAP0805",
            "SMD Ceramic Capacitor", 2, "C2012X5R1H475K125AB", 0, 0
        ])
        # N3: GRM188R71H104KA93D (38 per board)
        ws_bom.append([
            "N3", "100nF", "Capacitor 0603 100nF 50VDC 10%", "C3..C40", "SMDCAP0603",
            "SMD Ceramic Capacitor", 38, "GRM188R71H104KA93D", 0, 0
        ])
        # N4: BLANK_LINK_PART (1 per board)
        ws_bom.append([
            "N4", "10uF", "Blank link cap", "C5", "0805",
            "Capacitor", 1, "BLANK_LINK_PART", 0, 0
        ])
        # N5: UNSUPPORTED_PART (1 per board)
        ws_bom.append([
            "N5", "IC", "Unsupported store", "U1", "SOIC-8",
            "IC", 1, "UNSUPPORTED_PART", 0, 0
        ])
        # N6: UNRESOLVED_PART_NOT_IN_DB (1 per board, not present in inventory)
        ws_bom.append([
            "N6", "MCU", "Mystery MCU not in DB", "U2", "LQFP-64",
            "Microcontroller", 1, "MYSTERY_MCU_999", 0, 0
        ])
        bom_wb.save(bom_path)
        bom_wb.close()

        return inv_path, bom_path

    def test_path_b_bom_import_and_quantities_preserved(self, bom_and_inventory_path_b, dual_path_registry):
        """BOM import preserves unit quantities and multiplier scaling (N = unit_qty * M)."""
        inv_path, bom_path = bom_and_inventory_path_b
        loader = ExcelDataLoader()
        loader.load_file(inv_path)

        board_multiplier = 5
        ok, msg, result = BOMAuditor.audit_bom(
            bom_path, inventory=loader, board_multiplier=board_multiplier, mode=ProcurementMode.DATABASE_AWARE
        )
        assert ok is True
        assert result is not None
        assert result.total_items_count == 6

        # Check TAJD107K016RNJ: 1 * 5 = 5 needed; in stock = 1 -> shortage = 4
        taj = next(it for it in result.items if it.part_number == "TAJD107K016RNJ")
        assert taj.unit_qty == 1
        assert taj.board_multiplier == 5
        assert taj.total_needed_qty == 5
        assert taj.db_stock_qty == 1
        assert taj.shortage_qty == 4
        assert taj.status == "SHORTAGE"

        # Check GRM188R71H104KA93D: 38 * 5 = 190 needed; in stock = 50 -> shortage = 140
        grm = next(it for it in result.items if it.part_number == "GRM188R71H104KA93D")
        assert grm.total_needed_qty == 190
        assert grm.db_stock_qty == 50
        assert grm.shortage_qty == 140

    def test_path_b_bom_resolves_inventory_record_and_supplier_links(self, bom_and_inventory_path_b, dual_path_registry):
        """BOM component resolves to matching inventory record and parses supplier-link mapping."""
        inv_path, bom_path = bom_and_inventory_path_b
        loader = ExcelDataLoader()
        loader.load_file(inv_path)

        ok, msg, result = BOMAuditor.audit_bom(bom_path, inventory=loader, board_multiplier=1)
        assert ok is True

        taj = next(it for it in result.items if it.part_number == "TAJD107K016RNJ")
        assert taj.is_multi_vendor is True
        assert taj.is_orderable is True
        assert taj.link_type == "MULTI"
        assert len(taj.vendor_links) == 2
        assert "vendor_lion" in taj.vendor_links
        assert "vendor_javan" in taj.vendor_links

        c2012 = next(it for it in result.items if it.part_number == "C2012X5R1H475K125AB")
        assert c2012.is_multi_vendor is False
        assert c2012.is_orderable is True
        assert c2012.link_type == "JAVAN"
        assert c2012.vendor_id == "vendor_javan"

        grm = next(it for it in result.items if it.part_number == "GRM188R71H104KA93D")
        assert grm.is_multi_vendor is False
        assert grm.is_orderable is True
        assert grm.link_type == "LION"
        assert grm.vendor_id == "vendor_lion"

    def test_path_b_single_link_bom_item_bypasses_dialog(self, qapp, bom_and_inventory_path_b, dual_path_registry):
        """Single-link BOM shortages must bypass VendorSelectionDialog completely."""
        inv_path, bom_path = bom_and_inventory_path_b
        loader = ExcelDataLoader()
        loader.load_file(inv_path)

        ok, msg, result = BOMAuditor.audit_bom(bom_path, inventory=loader, board_multiplier=1)
        single_shortages = [it.to_dict() for it in result.items if it.part_number in ("C2012X5R1H475K125AB", "GRM188R71H104KA93D")]

        dlg = OrderPreviewDialog(all_shortages=single_shortages, registry=dual_path_registry)
        assert len(dlg.orderable_items) == 2
        assert len(dlg.multi_vendor_items) == 0

        with patch.object(VendorSelectionDialog, "exec") as mock_dlg_exec:
            with patch("ui.dialogs.order_preview_dialog.OrderWorkerThread") as mock_worker_cls:
                with patch("ui.dialogs.order_preview_dialog.settings.get", return_value="mock_user"):
                    dlg._execute_mode_a_order()
                    mock_dlg_exec.assert_not_called()
                    assert mock_worker_cls.called

    def test_path_b_multi_link_bom_item_prompts_and_dispatches_lion(self, qapp, bom_and_inventory_path_b, dual_path_registry):
        """Multi-vendor BOM shortage prompts user and dispatches only Lion URL when selected."""
        inv_path, bom_path = bom_and_inventory_path_b
        loader = ExcelDataLoader()
        loader.load_file(inv_path)

        ok, msg, result = BOMAuditor.audit_bom(bom_path, inventory=loader, board_multiplier=5)
        taj_shortage = [it.to_dict() for it in result.items if it.part_number == "TAJD107K016RNJ"]

        dlg = OrderPreviewDialog(all_shortages=taj_shortage, registry=dual_path_registry)
        assert len(dlg.multi_vendor_items) == 1

        def select_lion(*args, **kwargs):
            dlg.multi_vendor_items[0]["selected_vendor_id"] = "vendor_lion"
            dlg.multi_vendor_items[0]["selected_url"] = dlg.multi_vendor_items[0]["vendor_links"]["vendor_lion"]
            return QDialog.DialogCode.Accepted

        with patch.object(VendorSelectionDialog, "exec", side_effect=select_lion) as mock_exec:
            with patch("ui.dialogs.order_preview_dialog.OrderWorkerThread") as mock_worker_cls:
                with patch("ui.dialogs.order_preview_dialog.settings.get", return_value="mock_user"):
                    dlg._execute_mode_a_order()
                    assert mock_exec.called
                    assert mock_worker_cls.called

                    dispatched_items = mock_worker_cls.call_args[1]["items"]
                    req = ProcurementBridge.item_to_request(dispatched_items[0], registry=dual_path_registry)
                    assert req.explicit_vendor_id == "vendor_lion"
                    assert req.explicit_vendor_url == "https://lionelectronic.ir/products/3775-TAJD107K016RNJ"
                    assert req.required_quantity == 4
                    assert "javanelec" not in req.explicit_vendor_url

    def test_path_b_multi_link_bom_item_prompts_and_dispatches_javan(self, qapp, bom_and_inventory_path_b, dual_path_registry):
        """Multi-vendor BOM shortage prompts user and dispatches only Javan URL when selected."""
        inv_path, bom_path = bom_and_inventory_path_b
        loader = ExcelDataLoader()
        loader.load_file(inv_path)

        ok, msg, result = BOMAuditor.audit_bom(bom_path, inventory=loader, board_multiplier=5)
        taj_shortage = [it.to_dict() for it in result.items if it.part_number == "TAJD107K016RNJ"]

        dlg = OrderPreviewDialog(all_shortages=taj_shortage, registry=dual_path_registry)

        def select_javan(*args, **kwargs):
            dlg.multi_vendor_items[0]["selected_vendor_id"] = "vendor_javan"
            dlg.multi_vendor_items[0]["selected_url"] = dlg.multi_vendor_items[0]["vendor_links"]["vendor_javan"]
            return QDialog.DialogCode.Accepted

        with patch.object(VendorSelectionDialog, "exec", side_effect=select_javan) as mock_exec:
            with patch("ui.dialogs.order_preview_dialog.OrderWorkerThread") as mock_worker_cls:
                dlg._execute_mode_a_order()
                assert mock_exec.called
                assert mock_worker_cls.called

                dispatched_items = mock_worker_cls.call_args[1]["items"]
                req = ProcurementBridge.item_to_request(dispatched_items[0], registry=dual_path_registry)
                assert req.explicit_vendor_id == "vendor_javan"
                assert req.explicit_vendor_url == "https://www.javanelec.com/shop/product/55407/toshiba-corporation/t2n7002bk"
                assert "lionelectronic" not in req.explicit_vendor_url

    def test_path_b_cancellation_prevents_dispatch(self, qapp, bom_and_inventory_path_b, dual_path_registry):
        """Cancelling the dialog from the BOM path aborts procurement without starting worker."""
        inv_path, bom_path = bom_and_inventory_path_b
        loader = ExcelDataLoader()
        loader.load_file(inv_path)

        ok, msg, result = BOMAuditor.audit_bom(bom_path, inventory=loader, board_multiplier=1)
        taj_shortage = [it.to_dict() for it in result.items if it.part_number == "TAJD107K016RNJ"]

        dlg = OrderPreviewDialog(all_shortages=taj_shortage, registry=dual_path_registry)

        with patch.object(VendorSelectionDialog, "exec", return_value=QDialog.DialogCode.Rejected) as mock_exec:
            with patch("ui.dialogs.order_preview_dialog.OrderWorkerThread") as mock_worker_cls:
                dlg._execute_mode_a_order()
                assert mock_exec.called
                mock_worker_cls.assert_not_called()
                assert dlg.worker is None

    def test_path_b_unresolved_or_missing_product_not_invented(self, bom_and_inventory_path_b, dual_path_registry):
        """An unresolved or missing product record (NOT_IN_DB) never receives an invented URL."""
        inv_path, bom_path = bom_and_inventory_path_b
        loader = ExcelDataLoader()
        loader.load_file(inv_path)

        ok, msg, result = BOMAuditor.audit_bom(bom_path, inventory=loader, board_multiplier=1)
        unresolved = next(it for it in result.items if it.part_number == "MYSTERY_MCU_999")
        assert unresolved.status == "NOT_IN_DB"
        assert unresolved.link == ""
        assert unresolved.is_orderable is False
        assert unresolved.vendor_id is None
        assert unresolved.vendor_links == {}

        # Must not appear in orderable shortages
        orderable = result.get_orderable_shortages()
        assert not any(s["part_number"] == "MYSTERY_MCU_999" for s in orderable)

    def test_path_b_missing_or_invalid_link_produces_explicit_unavailable(self, bom_and_inventory_path_b, dual_path_registry):
        """BOM items with blank or unsupported links produce explicit unavailable results."""
        inv_path, bom_path = bom_and_inventory_path_b
        loader = ExcelDataLoader()
        loader.load_file(inv_path)

        ok, msg, result = BOMAuditor.audit_bom(bom_path, inventory=loader, board_multiplier=1)

        # Blank store link part
        blank_it = next(it for it in result.items if it.part_number == "BLANK_LINK_PART")
        assert blank_it.status == "SHORTAGE"
        assert blank_it.link == ""
        assert blank_it.is_orderable is False
        assert blank_it.link_type == "NONE"

        # Unsupported store link part
        unsupported_it = next(it for it in result.items if it.part_number == "UNSUPPORTED_PART")
        assert unsupported_it.status == "SHORTAGE"
        assert unsupported_it.is_orderable is False
        assert unsupported_it.link_type == "INVALID"
        assert unsupported_it.link_domain == "unsupportedstore.com"

        # Check unsupported list
        unsupported = result.get_unsupported_shortages()
        unsupported_pns = [s["part_number"] for s in unsupported]
        assert "BLANK_LINK_PART" in unsupported_pns
        assert "UNSUPPORTED_PART" in unsupported_pns


# =============================================================================
# SHARED REGRESSION & CONTRACT PRESERVATION TESTS
# =============================================================================

class TestSharedModeAContracts:
    """
    Validates shared Mode A contracts and invariants across both entry points.
    """

    def test_shared_contract_legacy_single_url_compatibility(self, dual_path_registry):
        """Legacy plain URL values resolve correctly without requiring LABEL= syntax."""
        url = "https://lionelectronic.ir/products/2070-NE555P"
        parsed = parse_vendor_links(url, registry=dual_path_registry)
        assert parsed.is_valid is True
        assert parsed.count == 1
        entry = parsed.get_single_entry()
        assert entry.vendor_label == "LION"
        assert entry.vendor_id == "vendor_lion"
        assert entry.url == url

    def test_shared_contract_raw_combined_string_never_dispatched(self, dual_path_registry):
        """Combined multi-vendor string must never be passed to resolve_adapter_for_url."""
        raw_multi = (
            "LION=https://lionelectronic.ir/products/3775 || "
            "JAVAN=https://www.javanelec.com/shop/product/55407"
        )
        vid, adapter = dual_path_registry.resolve_adapter_for_url(raw_multi)
        assert vid is None
        assert adapter is None

    def test_unsupported_procurement_mode_rejected(self, dual_path_registry):
        """Unsupported modes are rejected and Mode A requires registered links."""
        item_unsupported = {
            "mode": "UNSUPPORTED_MODE",
            "part_number": "STM32F401RET6",
            "needed_qty": 5,
            "link": "",
        }
        is_ord, vid = ProcurementBridge.is_item_orderable(item_unsupported, dual_path_registry)
        assert is_ord is False
        assert vid is None

        with pytest.raises(ValueError, match="sole supported procurement workflow"):
            ProcurementBridge.item_to_request(item_unsupported)

        item_mode_a = {
            "mode": ProcurementMode.DATABASE_AWARE,
            "part_number": "STM32F401RET6",
            "needed_qty": 5,
            "link": "",
        }
        is_ord_a, vid_a = ProcurementBridge.is_item_orderable(item_mode_a, dual_path_registry)
        assert is_ord_a is False
        assert vid_a is None
