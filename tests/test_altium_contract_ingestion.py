"""
Comprehensive Unit Tests for GAP-16: Extended Altium Parameter Contract Ingestion
Governing Document: DOC-OMNIX-ALTIUM-CONTRACT-001 (20-Parameter Canonical Contract)
"""

import os
import openpyxl
import pytest
from tempfile import NamedTemporaryFile

from bom_auditor import BOMAuditor, BOMItem
from core.procurement.models import (
    ProcurementMode,
    ProcurementRequest,
    ManufacturerRule,
    ManufacturerDataState,
)
from core.procurement.bridge import ProcurementBridge


@pytest.fixture
def full_20_param_bom_file():
    """Generates an Excel BOM containing all 20 canonical Altium contract parameters."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Altium_Contract_BOM"

    headers = [
        "Designator",
        "Quantity",
        "Comment",
        "Canonical_Component_ID",
        "MPN",
        "Manufacturer",
        "Approved_Manufacturers",
        "Manufacturer_Rule",
        "Package",
        "Package_Family",
        "Footprint_ID",
        "Pin_Count",
        "Description",
        "Value",
        "Tolerance",
        "Voltage_Rating",
        "Current_Rating",
        "Power_Rating",
        "Lifecycle_Status",
        "Library_Version",
    ]
    ws.append(headers)

    # Row 1: Fully qualified resistor
    ws.append([
        "R1, R2",
        "2",
        "10k Resistor",
        "RES-0805-10K-1%",
        "RC0805FR-0710KL",
        "Yageo",
        "Yageo, Vishay, Panasonic",
        "APPROVED_LIST",
        "0805",
        "SMD_CHIP",
        "RESC2012X60N",
        "2",
        "RES SMD 10K OHM 1% 1/8W 0805",
        "10k",
        "1%",
        "150V",
        "100mA",
        "0.125W",
        "Active",
        "v2.1",
    ])

    # Row 2: Microcontroller with EXACT rule
    ws.append([
        "U1",
        "1",
        "MCU 32-bit",
        "IC-MCU-STM32F103",
        "STM32F103C8T6",
        "STMicroelectronics",
        "",
        "EXACT",
        "LQFP-48",
        "QFP",
        "QFP50P900X900X160-48N",
        "48",
        "ARM Cortex-M3 72MHz 64KB Flash",
        "72MHz",
        "",
        "3.3V",
        "50mA",
        "",
        "Production",
        "v1.0",
    ])

    # Row 3: Generic Capacitor with ANY rule
    ws.append([
        "C1",
        "5.0",  # Float integer format
        "100nF Decoupling",
        "CAP-0603-100NF",
        "CC0603KRX7R9BB104",
        "ANY",
        "",
        "ANY",
        "0603",
        "SMD_CHIP",
        "CAPC1608X90N",
        "2",
        "CAP CER 0.1UF 50V X7R 0603",
        "100nF",
        "10%",
        "50V",
        "",
        "",
        "Active",
        "v2.0",
    ])

    with NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    wb.save(path)
    wb.close()
    yield path
    if os.path.exists(path):
        os.remove(path)


class TestAltiumContractIngestion:
    """Tests verifying 20-parameter ingestion, alias resolution, strict validation, and specs bundling."""

    def test_full_20_parameter_ingestion(self, full_20_param_bom_file):
        ok, msg, res = BOMAuditor.audit_bom(
            full_20_param_bom_file,
            inventory=None,
            board_multiplier=1,
            mode=ProcurementMode.DIRECT_BOM,
        )
        assert ok is True
        assert res is not None
        assert res.total_items_count == 3
        assert res.invalid_count == 0

        # Item 1: R1, R2
        r1 = res.items[0]
        assert r1.designator == "R1, R2"
        assert r1.unit_qty == 2
        assert r1.comment == "10k Resistor"
        assert r1.canonical_component_id == "RES-0805-10K-1%"
        assert r1.part_number == "RC0805FR-0710KL"
        assert r1.manufacturer == "Yageo"
        assert r1.approved_manufacturers == ["Yageo", "Vishay", "Panasonic"]
        assert r1.manufacturer_rule == ManufacturerRule.APPROVED_LIST
        assert r1.mfg_data_state == ManufacturerDataState.EXPLICIT_BRAND_RULE
        assert r1.allowed_manufacturers == ["Yageo", "Vishay", "Panasonic"]
        assert r1.package == "0805"
        assert r1.package_family == "SMD_CHIP"
        assert r1.footprint == "RESC2012X60N"
        assert r1.pin_count == 2
        assert r1.description == "RES SMD 10K OHM 1% 1/8W 0805"
        assert r1.value == "10k"
        assert r1.tolerance == "1%"
        assert r1.voltage_rating == "150V"
        assert r1.current_rating == "100mA"
        assert r1.power_rating == "0.125W"
        assert r1.lifecycle_status == "Active"
        assert r1.library_version == "v2.1"
        assert r1.required_specs == {
            "value": "10k",
            "tolerance": "1%",
            "voltage_rating": "150V",
            "current_rating": "100mA",
            "power_rating": "0.125W",
        }
        assert r1.required_dimensions == {}

        # Item 2: U1
        u1 = res.items[1]
        assert u1.part_number == "STM32F103C8T6"
        assert u1.manufacturer == "STMicroelectronics"
        assert u1.manufacturer_rule == ManufacturerRule.EXACT
        assert u1.mfg_data_state == ManufacturerDataState.EXPLICIT_BRAND_RULE
        assert u1.allowed_manufacturers == ["STMicroelectronics"]
        assert u1.is_manufacturer_any is False
        assert u1.pin_count == 48

        # Item 3: C1 (with ANY rule and float quantity "5.0")
        c1 = res.items[2]
        assert c1.unit_qty == 5
        assert c1.manufacturer_rule == ManufacturerRule.ANY
        assert c1.mfg_data_state == ManufacturerDataState.EXPLICIT_ANY
        assert c1.is_manufacturer_any is True
        assert c1.allowed_manufacturers is None

    def test_missing_mpn_never_falls_back_to_comment(self):
        """
        FORENSIC VERIFICATION:
        MPN missing / blank -> INVALID_REQUIRED_PARAMETER / NON_ORDERABLE
        Never use Comment as MPN.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Designator", "Quantity", "Comment", "MPN"])
        ws.append(["R1", "1", "10k 0805 Resistor", ""])  # MPN blank!
        ws.append(["R2", "1", "100nF 50V", "   "])  # MPN whitespace only!

        with NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        wb.save(path)
        wb.close()

        try:
            ok, msg, res = BOMAuditor.audit_bom(
                path,
                inventory=None,
                board_multiplier=1,
                mode=ProcurementMode.DIRECT_BOM,
            )
            assert ok is True
            assert res is not None
            assert res.total_items_count == 2
            assert res.invalid_count == 2

            for it in res.items:
                assert it.is_valid is False
                assert it.part_number == ""  # Did NOT copy comment!
                assert it.is_orderable is False
                assert any("INVALID_REQUIRED_PARAMETER" in err for err in it.validation_errors)

            # Check Bridge translation throws ValueError
            invalid_item_dict = res.items[0].to_dict()
            with pytest.raises(ValueError, match="part_number must be a non-empty string"):
                ProcurementBridge.item_to_request(invalid_item_dict)

        finally:
            if os.path.exists(path):
                os.remove(path)

    def test_strict_quantity_validation(self):
        """
        Tests validation for various quantity inputs:
        - Non-integer e.g. '2.5' -> REJECTED
        - Zero or negative e.g. '0', '-3' -> REJECTED
        - Non-numeric string e.g. 'N/A' -> REJECTED
        - Valid float integer e.g. '4.0' -> ACCEPTED as 4
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Designator", "Quantity", "MPN"])
        ws.append(["D1", "2.5", "1N4148"])    # Non-integer float
        ws.append(["D2", "0", "1N4007"])      # Zero
        ws.append(["D3", "-1", "BAT54S"])     # Negative
        ws.append(["D4", "abc", "SS14"])      # Non-numeric
        ws.append(["D5", "4.0", "SS34"])      # Valid float integer
        ws.append(["D6", "10", "MBR0520"])    # Valid standard int

        with NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        wb.save(path)
        wb.close()

        try:
            ok, msg, res = BOMAuditor.audit_bom(
                path,
                inventory=None,
                board_multiplier=1,
                mode=ProcurementMode.DIRECT_BOM,
            )
            assert ok is True
            assert res.total_items_count == 6
            assert res.invalid_count == 4

            # D1: Non-integer
            d1 = res.items[0]
            assert d1.is_valid is False
            assert any("INVALID_QUANTITY" in err for err in d1.validation_errors)

            # D2: Zero
            d2 = res.items[1]
            assert d2.is_valid is False
            assert any("INVALID_QUANTITY" in err for err in d2.validation_errors)

            # D3: Negative
            d3 = res.items[2]
            assert d3.is_valid is False
            assert any("INVALID_QUANTITY" in err for err in d3.validation_errors)

            # D4: Non-numeric
            d4 = res.items[3]
            assert d4.is_valid is False
            assert any("INVALID_QUANTITY" in err for err in d4.validation_errors)

            # D5: Float integer
            d5 = res.items[4]
            assert d5.is_valid is True
            assert d5.unit_qty == 4

            # D6: Standard int
            d6 = res.items[5]
            assert d6.is_valid is True
            assert d6.unit_qty == 10

        finally:
            if os.path.exists(path):
                os.remove(path)

    def test_persian_and_alternative_header_alias_resolution(self):
        """Verify that Persian and legacy headers resolve to the canonical contract."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["محل قطعه", "تعداد", "شماره فنی", "سازنده", "پکیج", "فوت پرینت", "شرح", "مقدار", "ولتاژ", "جریان", "توان"])
        ws.append(["U1", "1", "NE555P", "TI", "DIP-8", "DIP8", "تایمر دقیق", "555", "15V", "200mA", "600mW"])

        with NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        wb.save(path)
        wb.close()

        try:
            ok, msg, res = BOMAuditor.audit_bom(path, inventory=None, mode=ProcurementMode.DIRECT_BOM)
            assert ok is True
            assert res.total_items_count == 1
            it = res.items[0]
            assert it.designator == "U1"
            assert it.unit_qty == 1
            assert it.part_number == "NE555P"
            assert it.manufacturer == "TI"
            assert it.package == "DIP-8"
            assert it.footprint == "DIP8"
            assert it.description == "تایمر دقیق"
            assert it.required_specs["value"] == "555"
            assert it.required_specs["voltage_rating"] == "15V"
            assert it.required_specs["current_rating"] == "200mA"
            assert it.required_specs["power_rating"] == "600mW"
        finally:
            if os.path.exists(path):
                os.remove(path)
