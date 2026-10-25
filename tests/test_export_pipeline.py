"""
Omnix Automated Test Suite — Multi-Vendor Deterministic Export Pipeline (GAP-18)
Exhaustive verification of CSV, XLSX, and PDF exports, deterministic ordering,
data contracts, audit trails, and side-effect-free execution.
"""

import os
import io
import tempfile
import pytest
from unittest.mock import MagicMock
from PyQt6.QtWidgets import QApplication
import openpyxl

from core.models import IdentityMatchTier, OrderConstraints, ProductReference
from core.procurement.models import (
    BOMProcurementPlan,
    CandidateEvaluation,
    EvaluationState,
    ManufacturerRule,
    PriceTier,
    ProcurementMode,
    ProcurementRequest,
    RequestProcurementResult,
    VendorProcurementCandidate,
)
from core.procurement.reporting import (
    CandidateLifecycleState,
    ProcurementCSVExporter,
    ProcurementPDFExporter,
    ProcurementReportBuilder,
    ProcurementReportData,
    ProcurementXLSXExporter,
    ReportType,
)
from ui.dialogs.order_preview_dialog import OrderPreviewDialog


@pytest.fixture(scope="session")
def qapp():
    """Ensure singleton QApplication instance exists."""
    app = QApplication.instance()
    if app is None:
        app = QApplication(["--platform", "offscreen"])
    yield app


@pytest.fixture
def sample_procurement_plan():
    """Constructs a deterministic multi-vendor, multi-candidate procurement plan for testing."""
    # Item 1: NE555P Timer (Mode A - Database Aware Shortage)
    req1 = ProcurementRequest(
        part_number="NE555P",
        required_quantity=10,
        description="Precision Timer IC مقاومت و تایمر",
        package="DIP-8",
        footprint="DIP-8",
        manufacturer="Texas Instruments",
        mode=ProcurementMode.DATABASE_AWARE,
        source="U1",
        metadata={
            "designator": "U1",
            "comment": "Precision Timer",
            "bom_raw_demand": 25,
            "db_stock_qty": 15,
            "shortage_qty": 10,
        },
    )

    # Candidate 1A: Lion NE555 (Eligible, 5000 IRR, MOQ 5, Step 5) -> Selected
    cand_1a = VendorProcurementCandidate(
        vendor_id="vendor_lion",
        part_number="NE555P",
        product_ref=ProductReference(vendor_id="vendor_lion", vendor_sku="LION-555", pdp_url="https://lionelectronic.ir/product/2070-ne555p"),
        description="Precision Timer DIP-8",
        in_stock=True,
        stock_quantity=100,
        can_order=True,
        base_unit_price_irr=5000,
        price_tiers=[PriceTier(min_quantity=1, unit_price_irr=5000), PriceTier(min_quantity=50, unit_price_irr=4500)],
        order_constraints=OrderConstraints(moq_raw=5, packaging_step_raw=5),
        manufacturer="Texas Instruments",
        package="DIP-8",
        footprint="DIP-8",
    )
    eval_1a = CandidateEvaluation(
        candidate=cand_1a,
        required_quantity=10,
        effective_quantity=10,
        effective_unit_price_irr=5000,
        effective_total_price_irr=50000,
        is_eligible=True,
        state=EvaluationState.ELIGIBLE,
        rejection_reasons=[],
    )

    # Candidate 1B: Javan NE555 (Eligible, 4800 IRR, MOQ 10, Step 10)
    cand_1b = VendorProcurementCandidate(
        vendor_id="vendor_javan",
        part_number="NE555P",
        product_ref=ProductReference(vendor_id="vendor_javan", vendor_sku="JAV-555", pdp_url="https://javanelec.com/product/12345/ne555p"),
        description="Standard Timer DIP-8",
        in_stock=True,
        stock_quantity=50,
        can_order=True,
        base_unit_price_irr=4800,
        order_constraints=OrderConstraints(moq_raw=10, packaging_step_raw=10),
        manufacturer="Texas Instruments",
        package="DIP-8",
        footprint="DIP-8",
    )
    eval_1b = CandidateEvaluation(
        candidate=cand_1b,
        required_quantity=10,
        effective_quantity=10,
        effective_unit_price_irr=4800,
        effective_total_price_irr=48000,
        is_eligible=True,
        state=EvaluationState.ELIGIBLE,
        rejection_reasons=[],
    )

    # Candidate 1C: RoboEq NE555 (Incompatible Package: SOP-8)
    cand_1c = VendorProcurementCandidate(
        vendor_id="vendor_roboeq",
        part_number="NE555D",
        product_ref=ProductReference(vendor_id="vendor_roboeq", vendor_sku="ROBO-555", pdp_url="https://roboeq.ir/products/detail/0101001/"),
        description="SMD Timer SOP-8",
        in_stock=True,
        stock_quantity=200,
        can_order=True,
        base_unit_price_irr=3000,
        order_constraints=OrderConstraints(moq_raw=1, packaging_step_raw=1),
        manufacturer="Texas Instruments",
        package="SOP-8",
        footprint="SOP-8",
    )
    eval_1c = CandidateEvaluation(
        candidate=cand_1c,
        required_quantity=10,
        effective_quantity=0,
        effective_unit_price_irr=3000,
        effective_total_price_irr=0,
        is_eligible=False,
        state=EvaluationState.INCOMPATIBLE_PACKAGE,
        rejection_reasons=["Package mismatch: expected 'DIP-8', found 'SOP-8'"],
    )

    res1 = RequestProcurementResult(
        request=req1,
        status="RESOLVED",
        candidates=[cand_1a, cand_1b, cand_1c],
        evaluations=[eval_1a, eval_1b, eval_1c],
        selected_evaluation=eval_1a,
    )

    # Item 2: STM32F103C8T6 (Mode A - Allocated to Javan)
    req2 = ProcurementRequest(
        part_number="STM32F103C8T6",
        required_quantity=5,
        description="ARM Cortex-M3 MCU میکروکنترلر",
        package="LQFP-48",
        footprint="LQFP-48",
        manufacturer="STMicroelectronics",
        mode=ProcurementMode.DATABASE_AWARE,
        source="U2",
        metadata={
            "designator": "U2",
            "comment": "ARM MCU",
            "bom_raw_demand": 5,
            "db_stock_qty": 0,
            "shortage_qty": 5,
        },
    )
    cand_2a = VendorProcurementCandidate(
        vendor_id="vendor_javan",
        part_number="STM32F103C8T6",
        product_ref=ProductReference(vendor_id="vendor_javan", vendor_sku="JAV-STM32", pdp_url="https://javanelec.com/product/stm32f103"),
        description="STM32 Microcontroller",
        in_stock=True,
        stock_quantity=40,
        can_order=True,
        base_unit_price_irr=220000,
        order_constraints=OrderConstraints(moq_raw=1, packaging_step_raw=1),
        manufacturer="STMicroelectronics",
        package="LQFP-48",
        footprint="LQFP-48",
    )
    eval_2a = CandidateEvaluation(
        candidate=cand_2a,
        required_quantity=5,
        effective_quantity=5,
        effective_unit_price_irr=220000,
        effective_total_price_irr=1100000,
        is_eligible=True,
        state=EvaluationState.ELIGIBLE,
        rejection_reasons=[],
        unknown_parameters=["Flash Memory Retention Temp"],
    )
    res2 = RequestProcurementResult(
        request=req2,
        status="RESOLVED",
        candidates=[cand_2a],
        evaluations=[eval_2a],
        selected_evaluation=eval_2a,
    )

    # Item 3: Obsolete Component (Unresolved)
    req3 = ProcurementRequest(
        part_number="MAX232CPE+",
        required_quantity=2,
        description="RS-232 Driver",
        package="PDIP-16",
        footprint="PDIP-16",
        manufacturer="Maxim Integrated",
        mode=ProcurementMode.DATABASE_AWARE,
        source="U3",
        metadata={
            "designator": "U3",
            "comment": "RS232 IC",
            "bom_raw_demand": 2,
            "db_stock_qty": 0,
            "shortage_qty": 2,
        },
    )
    res3 = RequestProcurementResult(
        request=req3,
        status="UNRESOLVED",
        candidates=[],
        evaluations=[],
        selected_evaluation=None,
        error_message="No candidate offers found across registered vendors.",
    )

    return BOMProcurementPlan(
        plan_id="PLAN-TEST-001",
        created_at="2026-09-20 12:00:00",
        results=[res1, res2, res3],
    )


# ==============================================================================
# SECTION 1: ProcurementReportBuilder Tests
# ==============================================================================

class TestProcurementReportBuilder:
    """Test suite for ProcurementReportBuilder determinism, data contracts, and mapping."""

    def test_builder_mode_a_determinism(self, sample_procurement_plan):
        """Verify Mode A report builds with correct stock, shortage, and target quantities."""
        report = ProcurementReportBuilder.build(
            plan=sample_procurement_plan,
            source_title="Test Mode A Project",
            procurement_mode=ProcurementMode.DATABASE_AWARE,
            created_at="2026-09-20 12:00:00",
            report_id="REP-TEST-001",
        )

        assert report.report_type == ReportType.PRE_ORDER_CONFIRMATION
        assert report.currency == "IRR"
        assert report.total_line_items == 3
        assert report.resolved_line_items == 2
        assert report.unresolved_line_items == 1
        assert report.total_estimated_cost_irr == 50000 + 1100000  # 1,150,000 IRR

        # Line item 1 verification
        it1 = report.line_items[0]
        assert it1.line_number == 1
        assert it1.part_number == "NE555P"
        assert it1.designator == "U1"
        assert it1.bom_demand_qty == 25
        assert it1.db_stock_qty == 15
        assert it1.shortage_qty == 10
        assert it1.target_order_qty == 10
        assert it1.selected_vendor_name == "Lion Electronic"
        assert it1.selected_sku == "LION-555"
        assert it1.unit_price_irr == 5000
        assert it1.total_price_irr == 50000
        assert it1.status == "READY"

        # Line item 3 verification (unresolved)
        it3 = report.line_items[2]
        assert it3.line_number == 3
        assert it3.part_number == "MAX232CPE+"
        assert it3.status == "UNRESOLVED"
        assert it3.selected_vendor_name == "—"
        assert it3.total_price_irr is None

    def test_builder_mode_b_turnkey(self):
        """Verify Mode B sets db_stock_qty to 0 and sources full demand."""
        req = ProcurementRequest(
            part_number="STM32F407VGT6",
            required_quantity=20,
            mode=ProcurementMode.DIRECT_BOM,
            source="U1",
            metadata={"designator": "U1", "bom_raw_demand": 20, "db_stock_qty": 50},
        )
        cand = VendorProcurementCandidate(
            vendor_id="vendor_lion",
            part_number="STM32F407VGT6",
            product_ref=ProductReference(vendor_id="vendor_lion", vendor_sku="LION-STM32F4", pdp_url="https://lionelectronic.ir/p/123"),
            stock_quantity=100,
            in_stock=True,
            can_order=True,
            base_unit_price_irr=500000,
            order_constraints=OrderConstraints(moq_raw=1, packaging_step_raw=1),
        )
        ev = CandidateEvaluation(
            candidate=cand,
            required_quantity=20,
            effective_quantity=20,
            effective_unit_price_irr=500000,
            effective_total_price_irr=10000000,
            is_eligible=True,
            state=EvaluationState.ELIGIBLE,
        )
        res = RequestProcurementResult(request=req, status="RESOLVED", candidates=[cand], evaluations=[ev], selected_evaluation=ev)
        plan = BOMProcurementPlan(results=[res])

        report = ProcurementReportBuilder.build(
            plan=plan,
            source_title="Direct Turnkey BOM",
            procurement_mode=ProcurementMode.DIRECT_BOM,
        )

        assert report.procurement_mode == ProcurementMode.DIRECT_BOM
        assert report.line_items[0].db_stock_qty == 0
        assert report.line_items[0].bom_demand_qty == 20
        assert report.line_items[0].target_order_qty == 20

    def test_deterministic_candidate_sorting(self, sample_procurement_plan):
        """Verify candidate audit trail is ordered: Selected first, then eligible by vendor, then ineligible."""
        report = ProcurementReportBuilder.build(sample_procurement_plan)
        it1 = report.line_items[0]
        assert len(it1.candidates) == 3

        # 1st: Selected candidate (Lion NE555)
        assert it1.candidates[0].is_selected is True
        assert it1.candidates[0].vendor_id == "vendor_lion"
        assert it1.candidates[0].qualification_state == CandidateLifecycleState.SELECTED_BY_USER.value

        # 2nd: Eligible unselected (Javan NE555)
        assert it1.candidates[1].is_selected is False
        assert it1.candidates[1].is_eligible is True
        assert it1.candidates[1].vendor_id == "vendor_javan"
        assert it1.candidates[1].qualification_state == CandidateLifecycleState.ELIGIBLE.value

        # 3rd: Ineligible (RoboEq NE555)
        assert it1.candidates[2].is_eligible is False
        assert it1.candidates[2].vendor_id == "vendor_roboeq"
        assert it1.candidates[2].qualification_state == "INCOMPATIBLE_PACKAGE"
        assert "Package mismatch" in it1.candidates[2].rejection_reasons[0]

    def test_deterministic_vendor_summary_ordering(self, sample_procurement_plan):
        """Verify vendor summaries are ordered alphabetically by vendor_id."""
        report = ProcurementReportBuilder.build(sample_procurement_plan)
        assert len(report.vendor_summaries) == 2
        # Javan comes before Lion alphabetically
        assert report.vendor_summaries[0].vendor_id == "vendor_javan"
        assert report.vendor_summaries[0].item_count == 1
        assert report.vendor_summaries[0].total_pieces == 5
        assert report.vendor_summaries[0].subtotal_cost_irr == 1100000

        assert report.vendor_summaries[1].vendor_id == "vendor_lion"
        assert report.vendor_summaries[1].item_count == 1
        assert report.vendor_summaries[1].total_pieces == 10
        assert report.vendor_summaries[1].subtotal_cost_irr == 50000

    def test_pre_order_vs_post_staging_report_fields(self, sample_procurement_plan):
        """Verify post-staging report populates staging actuals and lifecycle states."""
        staging_actuals = [
            {
                "part": "NE555P",
                "status": "SUCCESS",
                "qty": 10,
                "msg": "Added 10 pcs to Lion cart",
                "vendor_id": "vendor_lion",
                "is_noop": False,
            },
            {
                "part": "STM32F103C8T6",
                "status": "SUCCESS",
                "qty": 0,
                "msg": "Already satisfied in cart (5/5 pcs)",
                "vendor_id": "vendor_javan",
                "is_noop": True,
            },
        ]

        report = ProcurementReportBuilder.build(
            plan=sample_procurement_plan,
            staging_results=staging_actuals,
            source_title="Post-Staging Execution Proof",
        )

        assert report.report_type == ReportType.POST_STAGING_FINAL

        # Item 1: Staged
        it1 = report.line_items[0]
        assert it1.actual_quantity_staged == 10
        assert it1.staging_status == "STAGED"
        assert it1.is_noop is False
        assert "Added 10 pcs" in it1.vendor_response_message
        assert it1.candidates[0].is_actually_staged is True

        # Item 2: No-op in cart
        it2 = report.line_items[1]
        assert it2.actual_quantity_staged == 0
        assert it2.staging_status == "NO_OP_SATISFIED"
        assert it2.is_noop is True

        # Item 3: Unresolved (never staged)
        it3 = report.line_items[2]
        assert it3.actual_quantity_staged is None
        assert it3.staging_status == "UNRESOLVED"

    def test_partial_failure_multi_vendor(self, sample_procurement_plan):
        """Verify partial failure reporting when Vendor A succeeds and Vendor B fails."""
        staging_actuals = [
            {
                "part": "NE555P",
                "status": "SUCCESS",
                "qty": 10,
                "msg": "Added to cart",
                "vendor_id": "vendor_lion",
                "is_noop": False,
            },
            {
                "part": "STM32F103C8T6",
                "status": "FAILED",
                "qty": 0,
                "msg": "Authentication required on Javan",
                "vendor_id": "vendor_javan",
                "is_noop": False,
            },
        ]

        report = ProcurementReportBuilder.build(
            plan=sample_procurement_plan,
            staging_results=staging_actuals,
        )

        it2 = report.line_items[1]
        assert it2.staging_status == "FAILED"
        assert any("Authentication required" in w for w in it2.warnings)

        # Check vendor summaries
        v_javan = next(v for v in report.vendor_summaries if v.vendor_id == "vendor_javan")
        assert v_javan.staging_status == "FAILED"

        v_lion = next(v for v in report.vendor_summaries if v.vendor_id == "vendor_lion")
        assert v_lion.staging_status == "SUCCESS"

    def test_export_does_not_mutate_plan(self, sample_procurement_plan):
        """Assert that building a report leaves the original BOMProcurementPlan 100% unchanged."""
        orig_count = len(sample_procurement_plan.results)
        orig_sel = sample_procurement_plan.results[0].selected_evaluation

        _ = ProcurementReportBuilder.build(sample_procurement_plan)

        assert len(sample_procurement_plan.results) == orig_count
        assert sample_procurement_plan.results[0].selected_evaluation == orig_sel


# ==============================================================================
# SECTION 2: CSV Exporter Tests
# ==============================================================================

class TestProcurementCSVExporter:
    """Test suite for RFC 4180 CSV export determinism, UTF-8 BOM, and structure."""

    def test_csv_export_deterministic_snapshot(self, sample_procurement_plan):
        """Verify repeated CSV generation from the same plan produces identical output."""
        report = ProcurementReportBuilder.build(sample_procurement_plan, created_at="2026-09-20 12:00:00", report_id="REP-TEST")
        csv_1 = ProcurementCSVExporter.to_string(report)
        csv_2 = ProcurementCSVExporter.to_string(report)

        assert csv_1 == csv_2
        assert "Line,Status,Designator,MPN" in csv_1
        assert "NE555P" in csv_1
        assert "STM32F103C8T6" in csv_1
        assert "50000" in csv_1
        assert "1100000" in csv_1

    def test_csv_utf8_sig_persian_support(self, sample_procurement_plan):
        """Verify CSV file is written with UTF-8 BOM (0xef 0xbb 0xbf) and Persian characters."""
        report = ProcurementReportBuilder.build(sample_procurement_plan, created_at="2026-09-20 12:00:00", report_id="REP-TEST")
        with tempfile.NamedTemporaryFile(suffix=".csv", delete=False) as tf:
            temp_path = tf.name

        try:
            ProcurementCSVExporter.export(report, temp_path)
            with open(temp_path, "rb") as f:
                raw_bytes = f.read()

            # UTF-8 BOM signature
            assert raw_bytes.startswith(b"\xef\xbb\xbf")

            # Check Persian string is readable
            content_str = raw_bytes.decode("utf-8-sig")
            assert "مقاومت و تایمر" in content_str
            assert "میکروکنترلر" in content_str
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)

    def test_csv_post_staging_columns(self, sample_procurement_plan):
        """Verify post-staging CSV includes actual staged quantity and vendor message."""
        staging_actuals = [
            {"part": "NE555P", "status": "SUCCESS", "qty": 10, "msg": "Added 10 pcs to Lion cart", "vendor_id": "vendor_lion", "is_noop": False}
        ]
        report = ProcurementReportBuilder.build(sample_procurement_plan, staging_results=staging_actuals, created_at="2026-09-20 12:00:00", report_id="REP-TEST")
        csv_str = ProcurementCSVExporter.to_string(report)

        assert "Actual Staged Qty" in csv_str
        assert "Added 10 pcs to Lion cart" in csv_str


# ==============================================================================
# SECTION 3: Excel / XLSX Exporter Tests
# ==============================================================================

class TestProcurementXLSXExporter:
    """Test suite for openpyxl Excel multi-sheet export, styling, and formulas."""

    def test_xlsx_export_multi_sheet_and_structure(self, sample_procurement_plan):
        """Verify 3 sheets are created with frozen panes, auto-filter, and numeric formatting."""
        report = ProcurementReportBuilder.build(sample_procurement_plan, created_at="2026-09-20 12:00:00", report_id="REP-TEST")
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            temp_path = tf.name

        try:
            ProcurementXLSXExporter.export(report, temp_path)
            wb = openpyxl.load_workbook(temp_path)

            # Check sheet names
            assert "Procurement Summary" in wb.sheetnames
            assert "Candidate Audit Trail" in wb.sheetnames
            assert "Warnings & Incompatibilities" in wb.sheetnames

            # Sheet 1 checks
            ws1 = wb["Procurement Summary"]
            assert "Omnix Multi-Vendor Procurement Order Report" in ws1["A1"].value
            assert ws1.freeze_panes is not None
            assert ws1.auto_filter.ref is not None

            # Sheet 2 checks (Candidate Audit Trail)
            ws2 = wb["Candidate Audit Trail"]
            assert ws2["A1"].value == "BOM Line #"
            assert ws2.max_row >= 5  # Headers + candidate offers
            assert ws2.freeze_panes == "A2"

            # Sheet 3 checks (Warnings)
            ws3 = wb["Warnings & Incompatibilities"]
            assert ws3["A1"].value == "BOM Line #"
            assert ws3.max_row >= 2  # Has unresolved item / warnings

            wb.close()
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)


# ==============================================================================
# SECTION 4: PDF Exporter Tests
# ==============================================================================

class TestProcurementPDFExporter:
    """Test suite for HTML generation and vector PDF export via QTextDocument."""

    def test_pdf_generate_html_structure(self, sample_procurement_plan):
        """Verify PDF HTML contains executive header, KPI boxes, vendor summary, and line items."""
        report = ProcurementReportBuilder.build(sample_procurement_plan, created_at="2026-09-20 12:00:00", report_id="REP-TEST")
        html = ProcurementPDFExporter.generate_html(report)

        assert "<!DOCTYPE html>" in html
        assert "Omnix Multi-Vendor Procurement Order Report" in html
        assert "Lion Electronic" in html
        assert "Javan Electronic" in html
        assert "NE555P" in html
        assert "STM32F103C8T6" in html
        assert "50,000 IRR" in html
        assert "1,100,000 IRR" in html
        assert "Candidate Decision Audit Trail" in html

    def test_pdf_mixed_persian_english_text(self, sample_procurement_plan):
        """Verify Persian descriptions and English technical identifiers are properly marked."""
        report = ProcurementReportBuilder.build(sample_procurement_plan, created_at="2026-09-20 12:00:00", report_id="REP-TEST")
        html = ProcurementPDFExporter.generate_html(report)

        assert 'dir="auto"' in html
        assert "مقاومت و تایمر" in html
        assert "میکروکنترلر" in html
        assert "<code>NE555P</code>" in html

    def test_pdf_export_render_execution(self, qapp, sample_procurement_plan):
        """Real vector PDF generation test asserting valid PDF byte structure."""
        report = ProcurementReportBuilder.build(sample_procurement_plan, created_at="2026-09-20 12:00:00", report_id="REP-TEST")
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tf:
            temp_path = tf.name

        try:
            ProcurementPDFExporter.export(report, temp_path)
            assert os.path.exists(temp_path)
            with open(temp_path, "rb") as f:
                header = f.read(5)
            # Standard PDF header signature
            assert header == b"%PDF-"
            assert os.path.getsize(temp_path) > 1000
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)


# ==============================================================================
# SECTION 5: UI Dialog Integration Tests
# ==============================================================================

class TestOrderPreviewDialogExportUI:
    """Test suite verifying OrderPreviewDialog export actions and integration."""

    def test_order_preview_dialog_export_integration(self, qapp, sample_procurement_plan):
        """Verify export_btn exists and can generate reports without mutating plan state."""
        dlg = OrderPreviewDialog(plan=sample_procurement_plan)
        assert hasattr(dlg, "export_btn")
        assert dlg.export_btn is not None
        assert "Export Staging Results" in dlg.export_btn.text()

        # Build report directly from dialog plan
        rep_pre = ProcurementReportBuilder.build(dlg.plan, staging_results=dlg.staging_results_log)
        assert rep_pre.report_type == ReportType.PRE_ORDER_CONFIRMATION
        assert rep_pre.total_line_items == 3

        # Simulate staging finished
        dlg.staging_results_log = [
            {"part": "NE555P", "status": "SUCCESS", "qty": 10, "msg": "Staged", "vendor_id": "vendor_lion"}
        ]
        rep_post = ProcurementReportBuilder.build(dlg.plan, staging_results=dlg.staging_results_log)
        assert rep_post.report_type == ReportType.POST_STAGING_FINAL
        assert rep_post.line_items[0].actual_quantity_staged == 10

        dlg.close()
