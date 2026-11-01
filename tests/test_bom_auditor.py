import os
import pytest
import openpyxl
from bom_auditor import BOMAuditor, normalize_string
from excel_loader import ExcelDataLoader

@pytest.fixture
def sample_inventory_and_bom(tmp_path):
    inv_path = os.path.join(str(tmp_path), 'Inventory.xlsx')
    inv_wb = openpyxl.Workbook()
    ws_inv = inv_wb.active
    ws_inv.title = 'Inventory'
    ws_inv.append(['ECODE', 'PART NAME', 'DISCRIPTION', 'Quantity in Stock', 'FOOTPRINT', 'PART NUMBER', 'LINK'])
    ws_inv.append([None, 'NE555P', 'Timer IC', 10, 'DIP-8', 'NE555P', 'https://lionelectronic.ir/products/2070-NE555P'])
    ws_inv.append([None, 'LM358N', 'Op-Amp', 5, 'DIP-8', 'LM358N', 'https://lionelectronic.ir/products/1932-LM358N'])
    ws_inv.append([None, '100uF 25V', 'Capacitor', 100, 'CAP-SMD-6.3', '100uF 25V', 'https://lionelectronic.ir/products/1000'])
    inv_wb.save(inv_path)
    inv_wb.close()

    bom_path = os.path.join(str(tmp_path), 'Sample_BOM.xlsx')
    bom_wb = openpyxl.Workbook()
    ws_bom = bom_wb.active
    ws_bom.append(['Designator', 'Quantity', 'Comment', 'Footprint', 'Description', 'Part Number'])
    ws_bom.append(['U1', 1, 'NE555P', 'DIP-8', 'Precision Timer', 'NE555P'])
    ws_bom.append(['U2, U3', 2, 'LM358N', 'DIP-8', 'Dual Op-Amp', 'LM358N'])
    ws_bom.append(['C1, C2, C3', 3, '100uF 25V', 'CAP-SMD-6.3', 'Capacitor', '100uF 25V'])
    ws_bom.append(['R1, R2', 2, '10k 0805', '0805', 'Resistor', 'RC0805-10K'])
    bom_wb.save(bom_path)
    bom_wb.close()

    return inv_path, bom_path

class TestBOMAuditor:
    @pytest.fixture(autouse=True)
    def register_test_adapters(self):
        from core.procurement.registry import VendorRegistry
        reg = VendorRegistry.get_default_registry()
        reg.register_default_adapters()
        yield
        reg.clear()

    def test_bom_audit_and_multiplier(self, sample_inventory_and_bom):
        inv_path, bom_path = sample_inventory_and_bom
        inv_loader = ExcelDataLoader()
        inv_loader.load_file(inv_path)

        ok, msg, res = BOMAuditor.audit_bom(bom_path, inv_loader, board_multiplier=10)
        assert ok is True
        assert res is not None
        assert res.total_items_count == 4
        assert res.board_multiplier == 10

        # 1. NE555P: 1 per board * 10 boards = 10 needed. In stock = 10 -> Shortage = 0 (IN_STOCK)
        ne555 = next(it for it in res.items if it.part_number == 'NE555P')
        assert ne555.total_needed_qty == 10
        assert ne555.db_stock_qty == 10
        assert ne555.shortage_qty == 0
        assert ne555.status == 'IN_STOCK'

        # 2. LM358N: 2 per board * 10 boards = 20 needed. In stock = 5 -> Shortage = 15 (SHORTAGE)
        lm358 = next(it for it in res.items if it.part_number == 'LM358N')
        assert lm358.total_needed_qty == 20
        assert lm358.db_stock_qty == 5
        assert lm358.shortage_qty == 15
        assert lm358.status == 'SHORTAGE'
        assert lm358.link_type == 'LION'

        # 3. 10k 0805: Not in DB -> NOT_IN_DB
        r_part = next(it for it in res.items if 'RC0805' in it.part_number)
        assert r_part.status == 'NOT_IN_DB'
        assert r_part.shortage_qty == 20
        assert r_part.total_needed_qty == 20

        # Test shortage lists
        all_shortages = res.get_all_shortages()
        orderable = res.get_orderable_shortages()
        assert len(all_shortages) == 1
        assert len(orderable) == 1
        assert orderable[0]['part_number'] == 'LM358N'
        assert orderable[0]['needed_qty'] == 15
        assert orderable[0]['is_orderable'] is True

    def test_bom_audit_multi_vendor_links(self, tmp_path):
        inv_path = os.path.join(str(tmp_path), 'Inventory_Multi.xlsx')
        inv_wb = openpyxl.Workbook()
        ws_inv = inv_wb.active
        ws_inv.title = 'Inventory'
        ws_inv.append(['ECODE', 'PART NAME', 'DISCRIPTION', 'Quantity in Stock', 'FOOTPRINT', 'PART NUMBER', 'LINK'])
        ws_inv.append([None, 'NE555P', 'Timer IC', 0, 'DIP-8', 'NE555P', 'https://lionelectronic.ir/products/2070-NE555P'])
        ws_inv.append([None, 'LM358N', 'Op-Amp', 0, 'DIP-8', 'LM358N', 'https://javanelec.com/product/123/LM358N'])
        ws_inv.append([None, 'ATmega328P', 'MCU', 0, 'DIP-28', 'ATmega328P', 'https://roboeq.ir/products/detail/0101001/'])
        ws_inv.append([None, 'CustomIC', 'Custom', 0, 'QFN-32', 'CustomIC', 'https://unknownstore.com/item/1'])
        inv_wb.save(inv_path)
        inv_wb.close()

        bom_path = os.path.join(str(tmp_path), 'Sample_BOM_Multi.xlsx')
        bom_wb = openpyxl.Workbook()
        ws_bom = bom_wb.active
        ws_bom.append(['Designator', 'Quantity', 'Comment', 'Footprint', 'Description', 'Part Number'])
        ws_bom.append(['U1', 1, 'NE555P', 'DIP-8', 'Timer', 'NE555P'])
        ws_bom.append(['U2', 1, 'LM358N', 'DIP-8', 'Op-Amp', 'LM358N'])
        ws_bom.append(['U3', 1, 'ATmega328P', 'DIP-28', 'MCU', 'ATmega328P'])
        ws_bom.append(['U4', 1, 'CustomIC', 'QFN-32', 'Custom', 'CustomIC'])
        bom_wb.save(bom_path)
        bom_wb.close()

        inv_loader = ExcelDataLoader()
        inv_loader.load_file(inv_path)

        ok, msg, res = BOMAuditor.audit_bom(bom_path, inv_loader, board_multiplier=1)
        assert ok is True
        assert res.total_items_count == 4

        shortages = res.get_all_shortages()
        assert len(shortages) == 4

        orderable = res.get_orderable_shortages()
        unsupported = res.get_unsupported_shortages()

        assert len(orderable) == 2  # Lion, Javan
        assert len(unsupported) == 2  # RoboEq (not installed) and unknownstore.com

        ne555 = next(s for s in orderable if s['part_number'] == 'NE555P')
        assert ne555['link_type'] == 'LION'
        assert ne555['vendor_id'] == 'vendor_lion'

        lm358 = next(s for s in orderable if s['part_number'] == 'LM358N')
        assert lm358['link_type'] == 'JAVAN'
        assert lm358['vendor_id'] == 'vendor_javan'

        custom = unsupported[0]
        assert custom['link_type'] == 'INVALID'
        assert custom['vendor_id'] is None
