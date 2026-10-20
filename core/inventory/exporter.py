"""
Omnix Core — Inventory Excel Exporter
Governing Document: Phase 2.4 Excel Export Specification

Exports the authoritative SQLite inventory database to a customer-facing Excel workbook
with EXACTLY two worksheets in exact order:
1. 'Inventory'      — Individual physical component records.
2. 'Assembly View'  — Aggregated stock groups and group-level policies.

Architectural guarantees:
- SQLite database remains the runtime source of truth.
- Output contains strictly 2 sheets (no extra reference/diagnostic sheets).
- Package leading zeroes are strictly preserved (e.g. '0402', '0603', '0805').
- Never modifies existing user workbooks.
"""

import os
from typing import Optional, List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.inventory.models import (
    InventoryProduct,
    StockPolicyGroup,
    normalize_package,
)
from core.inventory.repository import SqliteInventoryRepository


class InventoryExcelExporter:
    """
    Exports inventory and assembly stock groups to an Excel (.xlsx) workbook
    conforming strictly to the two-sheet contract: ['Inventory', 'Assembly View'].
    """

    @classmethod
    def export(
        cls,
        repository: SqliteInventoryRepository,
        filepath: str,
    ) -> str:
        """
        Exports the contents of the given SqliteInventoryRepository to an Excel workbook.

        Args:
            repository: SqliteInventoryRepository instance containing components and policies.
            filepath: Target output file path (.xlsx).

        Returns:
            The absolute path of the generated workbook.
        """
        wb = openpyxl.Workbook()
        # Remove default active sheet created by openpyxl
        default_sheet = wb.active

        # 1. Sheet 1: Inventory (Individual component records)
        ws_inventory = wb.create_sheet(title="Inventory")
        cls._build_inventory_sheet(ws_inventory, repository)

        # 2. Sheet 2: Assembly View (Aggregated stock groups & policies)
        ws_assembly = wb.create_sheet(title="Assembly View")
        cls._build_assembly_sheet(ws_assembly, repository)

        # Remove initial default sheet if still present
        if default_sheet is not None and default_sheet in wb.worksheets:
            wb.remove(default_sheet)

        # Ensure gridlines are displayed when opened in Excel
        try:
            if hasattr(ws_inventory, "views") and ws_inventory.views.sheetView:
                ws_inventory.views.sheetView[0].showGridLines = True
            if hasattr(ws_assembly, "views") and ws_assembly.views.sheetView:
                ws_assembly.views.sheetView[0].showGridLines = True
        except Exception:
            pass

        # Ensure directory exists and save
        out_dir = os.path.dirname(os.path.abspath(filepath))
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)
        wb.save(filepath)
        wb.close()
        return os.path.abspath(filepath)

    @classmethod
    def _build_inventory_sheet(cls, ws, repo: SqliteInventoryRepository) -> None:
        """Builds Sheet 1: 'Inventory' containing individual physical component records."""
        headers = [
            "Row",
            "Manufacturer",
            "Part Number (MPN)",
            "Category",
            "Type / Technology",
            "Value",
            "Package",
            "Footprint",
            "Quantity in Stock",
            "Description",
            "Purchase Link",
        ]

        font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        font_data = Font(name="Segoe UI", size=9, color="0F172A")
        font_code = Font(name="Consolas", size=9, color="0F172A")
        font_code_bold = Font(name="Consolas", size=9, bold=True, color="0F172A")
        font_number = Font(name="Consolas", size=9, color="0F172A")
        font_muted = Font(name="Consolas", size=9, color="64748B")
        font_link = Font(name="Segoe UI", size=9, color="0284C7", underline="single")

        thin_side = Side(border_style="thin", color="CBD5E1")
        border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        thick_bottom = Side(border_style="medium", color="0F172A")
        border_header = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thick_bottom)

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_wrap_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # Write Header
        ws.append(headers)
        ws.row_dimensions[1].height = 26
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border_header

        # Fetch all individual components
        components: List[InventoryProduct] = repo.get_all_products()

        for row_idx, comp in enumerate(components, start=2):
            norm_pkg = normalize_package(comp.package)
            qty = comp.quantity_in_stock if comp.quantity_in_stock is not None else 0
            is_odd = (row_idx % 2 == 1)

            row_data = [
                row_idx - 1,              # 1: Row
                comp.manufacturer,        # 2: Manufacturer
                comp.mpn,                 # 3: MPN
                comp.category,            # 4: Category
                comp.type_technology,     # 5: Type / Technology
                comp.value,               # 6: Value
                norm_pkg,                 # 7: Package (explicit string)
                comp.footprint,           # 8: Footprint
                qty,                      # 9: Quantity in Stock
                comp.description or "",   # 10: Description
                comp.purchase_link or "", # 11: Purchase Link
            ]
            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 21

            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = border_cell
                if is_odd:
                    cell.fill = fill_zebra

                if col_idx == 1:
                    cell.font = font_muted
                    cell.alignment = align_center
                elif col_idx in [2, 4, 5]:
                    cell.font = font_data
                    cell.alignment = align_left
                elif col_idx == 3:
                    cell.font = font_code_bold
                    cell.alignment = align_left
                elif col_idx in [6, 8]:
                    cell.font = font_code
                    cell.alignment = align_left
                elif col_idx == 7:
                    # Package: centered, Consolas font, explicit text format preserving leading zeroes
                    cell.font = font_code
                    cell.alignment = align_center
                    cell.number_format = "@"
                elif col_idx == 9:
                    cell.font = font_number
                    cell.alignment = align_center
                    cell.number_format = "#,##0"
                elif col_idx == 10:
                    cell.font = font_data
                    cell.alignment = align_wrap_left
                elif col_idx == 11:
                    cell.alignment = align_left
                    link_str = str(val or "").strip()
                    target_url = ""
                    if comp.supplier_links and comp.supplier_links[0].url:
                        target_url = comp.supplier_links[0].url.strip()
                    elif link_str:
                        if "=" in link_str:
                            candidate = link_str.split("=", 1)[1].split("||")[0].strip()
                            if candidate.startswith("http://") or candidate.startswith("https://"):
                                target_url = candidate
                        elif link_str.startswith("http://") or link_str.startswith("https://"):
                            target_url = link_str.split("||")[0].strip()

                    if target_url.startswith("http://") or target_url.startswith("https://"):
                        cell.font = font_link
                        cell.hyperlink = target_url
                    else:
                        cell.font = font_data

        # Auto-filter and responsive column widths
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{max(len(components) + 1, 1)}"
        cls._adjust_column_widths(ws, min_widths={
            1: 8,    # Row
            2: 20,   # Manufacturer
            3: 24,   # MPN
            4: 18,   # Category
            5: 22,   # Type
            6: 14,   # Value
            7: 12,   # Package
            8: 16,   # Footprint
            9: 18,   # Quantity in Stock
            10: 36,  # Description
            11: 34,  # Purchase Link
        })

    @classmethod
    def _build_assembly_sheet(cls, ws, repo: SqliteInventoryRepository) -> None:
        """Builds Sheet 2: 'Assembly View' containing aggregated stock groups and policies."""
        headers = [
            "Row",
            "Category",
            "Type / Technology",
            "Value",
            "Package",
            "Footprint",
            "Current Stock",
            "Min Stock",
            "Target Stock",
            "Policy Status",
            "Description",
        ]

        font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        fill_header = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        font_data = Font(name="Segoe UI", size=9, color="0F172A")
        font_code = Font(name="Consolas", size=9, color="0F172A")
        font_number = Font(name="Consolas", size=9, color="0F172A")
        font_muted = Font(name="Consolas", size=9, color="64748B")

        # Policy Status Fonts
        status_fonts = {
            "Adequate": Font(name="Segoe UI", size=9, bold=True, color="16A34A"),
            "Low Stock": Font(name="Segoe UI", size=9, bold=True, color="D97706"),
            "Out of Stock": Font(name="Segoe UI", size=9, bold=True, color="DC2626"),
            "Incomplete": Font(name="Segoe UI", size=9, bold=True, color="9333EA"),
        }

        thin_side = Side(border_style="thin", color="CBD5E1")
        border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        thick_bottom = Side(border_style="medium", color="0F172A")
        border_header = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thick_bottom)

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_wrap_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # Write Header
        ws.append(headers)
        ws.row_dimensions[1].height = 26
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border_header

        # Fetch aggregated groups from repository
        groups: List[StockPolicyGroup] = repo.get_all_groups()

        for row_idx, grp in enumerate(groups, start=2):
            norm_pkg = normalize_package(grp.package)
            curr_stock = grp.current_stock if grp.current_stock is not None else 0
            min_stk = grp.min_stock if grp.min_stock is not None else ""
            tgt_stk = grp.target_stock if grp.target_stock is not None else ""
            is_odd = (row_idx % 2 == 1)

            # Determine policy status string
            if not grp.is_complete:
                status_str = "Incomplete"
            elif grp.min_stock is not None and curr_stock < grp.min_stock:
                status_str = "Low Stock"
            elif curr_stock == 0:
                status_str = "Out of Stock"
            else:
                status_str = "Adequate"

            row_data = [
                row_idx - 1,             # 1: Row
                grp.category,            # 2: Category
                grp.type_technology,    # 3: Type / Technology
                grp.value,               # 4: Value
                norm_pkg,                # 5: Package (explicit string)
                grp.footprint,           # 6: Footprint
                curr_stock,              # 7: Current Stock
                min_stk,                 # 8: Min Stock
                tgt_stk,                 # 9: Target Stock
                status_str,              # 10: Policy Status
                grp.description or "",   # 11: Description
            ]
            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 21

            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = border_cell
                if is_odd:
                    cell.fill = fill_zebra

                if col_idx == 1:
                    cell.font = font_muted
                    cell.alignment = align_center
                elif col_idx in [2, 3]:
                    cell.font = font_data
                    cell.alignment = align_left
                elif col_idx in [4, 6]:
                    cell.font = font_code
                    cell.alignment = align_left
                elif col_idx == 5:
                    # Package: centered, Consolas font, explicit text format
                    cell.font = font_code
                    cell.alignment = align_center
                    cell.number_format = "@"
                elif col_idx in [7, 8, 9]:
                    cell.font = font_number
                    cell.alignment = align_center
                    if isinstance(val, (int, float)):
                        cell.number_format = "#,##0"
                    else:
                        cell.alignment = align_center
                elif col_idx == 10:
                    cell.font = status_fonts.get(status_str, font_data)
                    cell.alignment = align_center
                else:
                    cell.font = font_data
                    cell.alignment = align_wrap_left

        # Auto-filter and responsive column widths
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{max(len(groups) + 1, 1)}"
        cls._adjust_column_widths(ws, min_widths={
            1: 8,    # Row
            2: 18,   # Category
            3: 22,   # Type / Technology
            4: 14,   # Value
            5: 12,   # Package
            6: 16,   # Footprint
            7: 16,   # Current Stock
            8: 14,   # Min Stock
            9: 14,   # Target Stock
            10: 16,  # Policy Status
            11: 36,  # Description
        })

    @staticmethod
    def _adjust_column_widths(ws, min_widths: Dict[int, int]) -> None:
        """Autofit columns to content with defined minimum bounds."""
        for col_idx, col in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            col_letter = get_column_letter(col_idx)
            base_min = min_widths.get(col_idx, 10)
            ws.column_dimensions[col_letter].width = max(max_len + 3, base_min)
