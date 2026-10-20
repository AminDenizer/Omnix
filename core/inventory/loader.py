"""
Omnix Core - Inventory Database Loader
Governing Document: Phase 1 Master Architecture Specification
"""

import os
from typing import Optional, List, Dict, Any, Tuple
import openpyxl

from core.inventory.models import (
    InventoryProduct,
    StockPolicyGroup,
    ReferenceDataPayload,
    normalize_package,
    build_group_key,
)


class InventoryError(Exception):
    """Base exception for inventory database operations."""
    pass


class MissingWorksheetError(InventoryError):
    """Raised when an explicit required worksheet is absent from the workbook."""
    pass


class WorkbookAccessError(InventoryError):
    """Raised when the workbook file cannot be accessed or loaded."""
    pass


def resolve_inventory_database_path(configured_path: Optional[str] = None) -> str:
    """
    Resolve the inventory database path.
    Prioritizes explicit argument, then settings.get('inventory_database_path'),
    defaulting to 'autom/atom/Inventory.xlsm'.
    """
    path = configured_path
    if not path:
        try:
            from settings_manager import settings
            path = settings.get("inventory_database_path")
        except Exception:
            path = None

    if path and str(path).strip():
        p_str = str(path).strip()
        if os.path.isabs(p_str) and os.path.exists(p_str):
            return p_str
        p_cwd = os.path.abspath(p_str)
        if os.path.exists(p_cwd):
            return p_cwd
        try:
            from config import get_app_dir
            p_app = os.path.abspath(os.path.join(get_app_dir(), p_str))
            if os.path.exists(p_app):
                return p_app
        except Exception:
            pass
        return p_cwd

    # Default candidates
    candidates = [
        os.path.abspath("autom/atom/Inventory.xlsm"),
        os.path.join(os.getcwd(), "autom", "atom", "Inventory.xlsm"),
        os.path.join(os.getcwd(), "Inventory.xlsm"),
    ]
    try:
        from config import get_app_dir
        candidates.append(os.path.join(get_app_dir(), "Inventory.xlsm"))
        candidates.append(os.path.join(get_app_dir(), "autom", "atom", "Inventory.xlsm"))
    except Exception:
        pass

    for cand in candidates:
        if os.path.isfile(cand):
            return cand

    return os.path.abspath("autom/atom/Inventory.xlsm")


def _parse_int(val: Any) -> Optional[int]:
    """Parse an integer from cell value, returning None if empty or invalid."""
    if val is None or val == "":
        return None
    try:
        if isinstance(val, (int, float)):
            return int(val)
        s = str(val).strip()
        if not s:
            return None
        return int(float(s))
    except (ValueError, TypeError):
        return None


def _clean_str(val: Any) -> str:
    """Clean string from cell value, returning empty string if None."""
    if val is None:
        return ""
    return str(val).strip()


class InventoryDatabaseLoader:
    """
    Strictly read-only loader for the Omnix inventory database (Inventory.xlsm).

    Guarantees:
    - Never mutates or saves the source workbook.
    - Explicitly accesses named sheets: 'Stock Policy' and 'Inventory'.
    - Never relies on worksheet position or active worksheet.
    - Treats 'Stock Policy' as authoritative for group identity and policy values.
    - Reconciles individual Inventory products to Stock Policy groups using the 5 dimensions.
    - Preserves single purchase link and individual product identity.
    - Correctly normalizes package numbers (e.g. 603 -> '0603').
    - Reports missing or unusable formula values without inventing replacement aggregates.
    """

    SHEET_STOCK_POLICY = "Stock Policy"
    SHEET_INVENTORY = "Inventory"
    SHEET_REFERENCE_DATA = "Reference Data"

    def load_workbook(
        self,
        file_path: Optional[str] = None,
    ) -> Tuple[List[StockPolicyGroup], List[InventoryProduct], List[str]]:
        """
        Load inventory database from workbook.

        Args:
            file_path: Optional path to workbook. If omitted, resolved via resolve_inventory_database_path().

        Returns:
            Tuple of (policy_groups, inventory_products, warnings).

        Raises:
            WorkbookAccessError: If file not found or cannot be loaded.
            MissingWorksheetError: If 'Stock Policy' or 'Inventory' worksheet is absent.
        """
        resolved_path = resolve_inventory_database_path(file_path)

        if not os.path.exists(resolved_path):
            raise WorkbookAccessError(f"Inventory workbook not found at: {resolved_path}")

        warnings: List[str] = []

        try:
            wb = openpyxl.load_workbook(resolved_path, data_only=True, read_only=True)
        except Exception as e:
            raise WorkbookAccessError(f"Failed to open workbook '{resolved_path}': {e}") from e

        try:
            sheetnames = wb.sheetnames

            # 1. Verify explicitly named sheets
            if self.SHEET_STOCK_POLICY not in sheetnames:
                raise MissingWorksheetError(
                    f"Required worksheet '{self.SHEET_STOCK_POLICY}' missing from '{resolved_path}'. "
                    f"Available sheets: {sheetnames}"
                )
            if self.SHEET_INVENTORY not in sheetnames:
                raise MissingWorksheetError(
                    f"Required worksheet '{self.SHEET_INVENTORY}' missing from '{resolved_path}'. "
                    f"Available sheets: {sheetnames}"
                )

            # 2. Parse Stock Policy Sheet (Authoritative Group & Policy Source)
            ws_policy = wb[self.SHEET_STOCK_POLICY]
            policy_groups, group_map = self._parse_stock_policy_sheet(ws_policy, warnings)

            # 3. Parse Inventory Sheet (Individual Physical Product Records)
            ws_inv = wb[self.SHEET_INVENTORY]
            inventory_products = self._parse_inventory_sheet(ws_inv, warnings)

            # 4. Match individual Inventory products to Stock Policy groups
            matched_groups = self._match_products_to_groups(policy_groups, inventory_products, group_map)

            return matched_groups, inventory_products, warnings

        finally:
            try:
                wb.close()
            except Exception:
                pass

    def load_reference_data(
        self,
        file_path: Optional[str] = None,
    ) -> ReferenceDataPayload:
        """
        Load canonical reference data (manufacturers, categories, types) from the
        workbook's 'Reference Data' sheet, if present.

        Strictly read-only access.
        """
        resolved_path = resolve_inventory_database_path(file_path)
        if not os.path.exists(resolved_path):
            raise WorkbookAccessError(f"Inventory workbook not found at: {resolved_path}")

        try:
            wb = openpyxl.load_workbook(resolved_path, data_only=True, read_only=True)
        except Exception as e:
            raise WorkbookAccessError(f"Failed to open workbook '{resolved_path}': {e}") from e

        try:
            if self.SHEET_REFERENCE_DATA not in wb.sheetnames:
                return ReferenceDataPayload()

            ws = wb[self.SHEET_REFERENCE_DATA]
            rows_iter = ws.iter_rows(values_only=True)
            header = next(rows_iter, None)
            if not header:
                return ReferenceDataPayload()

            mfrs: List[str] = []
            seen_mfrs = set()
            cats: List[str] = []
            seen_cats = set()
            types: List[Tuple[str, str]] = []
            seen_types = set()

            for row in rows_iter:
                if not row:
                    continue
                # Column 0: Manufacturer
                if len(row) > 0 and row[0]:
                    m = str(row[0]).strip()
                    if m and m.lower() != "manufacturer" and m.upper() not in seen_mfrs:
                        seen_mfrs.add(m.upper())
                        mfrs.append(m)

                # Column 1 & 2: Category -> Type
                if len(row) > 1 and row[1]:
                    c = str(row[1]).strip()
                    if c and c.lower() != "category":
                        if c.upper() not in seen_cats:
                            seen_cats.add(c.upper())
                            cats.append(c)
                        if len(row) > 2 and row[2]:
                            t = str(row[2]).strip()
                            if t and t.lower() != "type/technology":
                                pair_key = (c.upper(), t.upper())
                                if pair_key not in seen_types:
                                    seen_types.add(pair_key)
                                    types.append((c, t))

                # Column 4: Category List
                if len(row) > 4 and row[4]:
                    c = str(row[4]).strip()
                    if c and c.lower() != "category list" and c.upper() not in seen_cats:
                        seen_cats.add(c.upper())
                        cats.append(c)

            # Category-specific columns in header (e.g. columns 23+)
            if header:
                for col_idx, col_name in enumerate(header):
                    if col_name:
                        col_str = str(col_name).strip()
                        if col_str.upper() in seen_cats:
                            canonical_cat = next((cat for cat in cats if cat.upper() == col_str.upper()), col_str)
                            for r in ws.iter_rows(values_only=True):
                                if len(r) > col_idx and r[col_idx]:
                                    t_val = str(r[col_idx]).strip()
                                    if t_val and t_val.upper() != col_str.upper():
                                        pair_key = (canonical_cat.upper(), t_val.upper())
                                        if pair_key not in seen_types:
                                            seen_types.add(pair_key)
                                            types.append((canonical_cat, t_val))

            return ReferenceDataPayload(
                manufacturers=tuple(mfrs),
                categories=tuple(cats),
                category_types=tuple(types),
            )
        finally:
            try:
                wb.close()
            except Exception:
                pass

    def _parse_stock_policy_sheet(
        self,
        ws: Any,
        warnings: List[str],
    ) -> Tuple[List[StockPolicyGroup], Dict[str, List[StockPolicyGroup]]]:
        """Parse the Stock Policy sheet and return policy groups and lookup map."""
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return [], {}

        col_map = self._map_columns(
            header_row,
            {
                "category": ["category", "دسته"],
                "type_technology": ["type/technology", "type", "technology", "نوع/تکنولوژی"],
                "value": ["value", "مقدار"],
                "package": ["package", "پکیج"],
                "footprint": ["footprint", "فوت پرینت"],
                "current_stock": ["total quantity in stock", "total stock", "موجودی کل"],
                "min_stock": ["min stock", "حداقل موجودی", "حد کسر"],
                "target_stock": ["target stock", "موجودی هدف", "تعداد سفارش"],
                "description": ["description", "شرح"],
            },
        )

        groups: List[StockPolicyGroup] = []
        group_map: Dict[str, List[StockPolicyGroup]] = {}

        for row_idx, row in enumerate(rows_iter, start=2):
            if not row:
                continue

            category = _clean_str(row[col_map["category"]]) if col_map["category"] is not None else ""
            if not category:
                # Empty row or formula placeholder evaluating to empty
                continue

            type_tech = _clean_str(row[col_map["type_technology"]]) if col_map["type_technology"] is not None else ""
            val = _clean_str(row[col_map["value"]]) if col_map["value"] is not None else ""
            raw_pkg = row[col_map["package"]] if col_map["package"] is not None else ""
            norm_pkg = normalize_package(raw_pkg)
            footprint = _clean_str(row[col_map["footprint"]]) if col_map["footprint"] is not None else ""
            desc = _clean_str(row[col_map["description"]]) if col_map["description"] is not None else ""

            # Policy threshold fields
            missing_fields: List[str] = []

            raw_cur_stock = row[col_map["current_stock"]] if col_map["current_stock"] is not None else None
            cur_stock = _parse_int(raw_cur_stock)
            if cur_stock is None:
                missing_fields.append("current_stock")
                warnings.append(
                    f"Row {row_idx} ({category}|{val}|{norm_pkg}): 'Total Quantity in Stock' "
                    f"cached formula result is missing or unusable ({raw_cur_stock!r})."
                )

            raw_min_stock = row[col_map["min_stock"]] if col_map["min_stock"] is not None else None
            min_stock = _parse_int(raw_min_stock)
            if min_stock is None:
                missing_fields.append("min_stock")
                warnings.append(
                    f"Row {row_idx} ({category}|{val}|{norm_pkg}): 'Min Stock' is missing or unusable ({raw_min_stock!r})."
                )

            raw_target_stock = row[col_map["target_stock"]] if col_map["target_stock"] is not None else None
            target_stock = _parse_int(raw_target_stock)
            if target_stock is None:
                missing_fields.append("target_stock")
                warnings.append(
                    f"Row {row_idx} ({category}|{val}|{norm_pkg}): 'Target Stock' is missing or unusable ({raw_target_stock!r})."
                )

            grp = StockPolicyGroup(
                category=category,
                type_technology=type_tech,
                value=val,
                package=norm_pkg,
                footprint=footprint,
                current_stock=cur_stock,
                min_stock=min_stock,
                target_stock=target_stock,
                description=desc,
                missing_fields=tuple(missing_fields),
            )
            groups.append(grp)
            k = grp.group_key
            if k not in group_map:
                group_map[k] = []
            group_map[k].append(grp)

        return groups, group_map

    def _parse_inventory_sheet(
        self,
        ws: Any,
        warnings: List[str],
    ) -> List[InventoryProduct]:
        """Parse individual physical product records from the Inventory sheet."""
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return []

        col_map = self._map_columns(
            header_row,
            {
                "manufacturer": ["manufacturer", "سازنده", "brand"],
                "mpn": ["mpn", "part number", "part_number", "شماره فنی"],
                "category": ["category", "دسته"],
                "type_technology": ["type/technology", "type", "technology", "نوع/تکنولوژی"],
                "package": ["package", "پکیج"],
                "footprint": ["footprint", "فوت پرینت"],
                "value": ["value", "مقدار"],
                "description": ["description", "شرح"],
                "quantity_in_stock": ["quantity in stock", "stock", "quantity", "موجودی"],
                "purchase_link": ["purchase link", "link", "url", "لینک خرید"],
            },
        )

        products: List[InventoryProduct] = []

        for row_idx, row in enumerate(rows_iter, start=2):
            if not row:
                continue

            mpn = _clean_str(row[col_map["mpn"]]) if col_map["mpn"] is not None else ""
            if not mpn:
                # Skip empty formatted table rows
                continue

            mfg = _clean_str(row[col_map["manufacturer"]]) if col_map["manufacturer"] is not None else ""
            cat = _clean_str(row[col_map["category"]]) if col_map["category"] is not None else ""
            type_tech = _clean_str(row[col_map["type_technology"]]) if col_map["type_technology"] is not None else ""
            raw_pkg_val = row[col_map["package"]] if col_map["package"] is not None else ""
            raw_pkg = str(raw_pkg_val).strip() if raw_pkg_val is not None else ""
            norm_pkg = normalize_package(raw_pkg_val)
            footprint = _clean_str(row[col_map["footprint"]]) if col_map["footprint"] is not None else ""
            val = _clean_str(row[col_map["value"]]) if col_map["value"] is not None else ""
            desc = _clean_str(row[col_map["description"]]) if col_map["description"] is not None else ""

            raw_qty = row[col_map["quantity_in_stock"]] if col_map["quantity_in_stock"] is not None else None
            qty = _parse_int(raw_qty)
            if qty is None:
                qty = 0

            raw_link = row[col_map["purchase_link"]] if col_map["purchase_link"] is not None else None
            link = str(raw_link).strip() if raw_link is not None and str(raw_link).strip() else None

            grp_k = build_group_key(cat, type_tech, val, norm_pkg, footprint)

            product = InventoryProduct(
                manufacturer=mfg,
                mpn=mpn,
                category=cat,
                type_technology=type_tech,
                package=norm_pkg,
                footprint=footprint,
                value=val,
                description=desc,
                quantity_in_stock=qty,
                purchase_link=link,
                group_key=grp_k,
                raw_package=raw_pkg,
            )
            products.append(product)

        return products

    def _match_products_to_groups(
        self,
        groups: List[StockPolicyGroup],
        products: List[InventoryProduct],
        group_map: Dict[str, List[StockPolicyGroup]],
    ) -> List[StockPolicyGroup]:
        """Associate matching physical products with their respective StockPolicyGroups."""
        products_by_key: Dict[str, List[InventoryProduct]] = {}
        for p in products:
            if p.group_key not in products_by_key:
                products_by_key[p.group_key] = []
            products_by_key[p.group_key].append(p)

        result_groups: List[StockPolicyGroup] = []
        for g in groups:
            matched_prods = tuple(products_by_key.get(g.group_key, []))
            # Create group with associated products
            hydrated_grp = StockPolicyGroup(
                category=g.category,
                type_technology=g.type_technology,
                value=g.value,
                package=g.package,
                footprint=g.footprint,
                current_stock=g.current_stock,
                min_stock=g.min_stock,
                target_stock=g.target_stock,
                description=g.description,
                products=matched_prods,
                missing_fields=g.missing_fields,
            )
            result_groups.append(hydrated_grp)

        return result_groups

    @staticmethod
    def _map_columns(
        header_row: Tuple[Any, ...],
        field_synonyms: Dict[str, List[str]],
    ) -> Dict[str, Optional[int]]:
        """Map canonical field names to column indexes using synonyms."""
        col_map: Dict[str, Optional[int]] = {k: None for k in field_synonyms}
        if not header_row:
            return col_map

        normalized_headers = [
            (idx, str(h).strip().lower()) for idx, h in enumerate(header_row) if h is not None
        ]

        for field_name, synonyms in field_synonyms.items():
            for idx, h_text in normalized_headers:
                for syn in synonyms:
                    if h_text == syn or syn in h_text:
                        col_map[field_name] = idx
                        break
                if col_map[field_name] is not None:
                    break

        return col_map
