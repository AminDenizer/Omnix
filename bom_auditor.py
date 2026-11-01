import os
import re
from typing import List, Dict, Any, Optional, Tuple
import openpyxl

from excel_loader import clean_cell_value, normalize_string, ExcelDataLoader
from core.procurement.models import (
    ProcurementMode,
    ManufacturerRule,
    ManufacturerDataState,
)
from core.logger import log_event, track_operation


class BOMItem:
    """
    Represents a single component line item from an Altium BOM export conforming to
    DOC-OMNIX-ALTIUM-CONTRACT-001 (Extended Altium 20-Parameter Canonical Contract).
    """

    def __init__(
        self,
        row_idx: int,
        designator: str = "",
        comment: str = "",
        description: str = "",
        footprint: str = "",
        lib_ref: str = "",
        part_number: str = "",
        unit_qty: int = 1,
        board_multiplier: int = 1,
        canonical_component_id: str = "",
        manufacturer: str = "",
        approved_manufacturers: Optional[List[str]] = None,
        manufacturer_rule: Optional[ManufacturerRule] = None,
        package: str = "",
        package_family: str = "",
        pin_count: Optional[int] = None,
        value: str = "",
        tolerance: str = "",
        voltage_rating: str = "",
        current_rating: str = "",
        power_rating: str = "",
        lifecycle_status: str = "",
        library_version: str = "",
        mode: ProcurementMode = ProcurementMode.DATABASE_AWARE,
        mfg_data_state: ManufacturerDataState = ManufacturerDataState.MISSING_MANUFACTURER_DATA,
        allowed_manufacturers: Optional[List[str]] = None,
        is_manufacturer_any: bool = False,
        required_specs: Optional[Dict[str, str]] = None,
        required_dimensions: Optional[Dict[str, str]] = None,
        is_valid: bool = True,
        validation_errors: Optional[List[str]] = None,
    ):
        self.row_idx = row_idx
        self.designator = designator
        self.comment = comment
        self.description = description
        self.footprint = footprint
        self.lib_ref = lib_ref
        self.part_number = part_number
        self.unit_qty = unit_qty
        self.board_multiplier = max(board_multiplier, 1)
        self.total_needed_qty = max(0, self.unit_qty * self.board_multiplier)

        # Extended Altium 20-Parameter Contract fields
        self.canonical_component_id = canonical_component_id
        self.manufacturer = manufacturer
        self.approved_manufacturers = approved_manufacturers or []
        self.manufacturer_rule = manufacturer_rule
        self.package = package
        self.package_family = package_family
        self.pin_count = pin_count
        self.value = value
        self.tolerance = tolerance
        self.voltage_rating = voltage_rating
        self.current_rating = current_rating
        self.power_rating = power_rating
        self.lifecycle_status = lifecycle_status
        self.library_version = library_version

        # Mode & Qualification state
        self.mode = mode if isinstance(mode, ProcurementMode) else ProcurementMode(mode)
        self.mfg_data_state = (
            mfg_data_state
            if isinstance(mfg_data_state, ManufacturerDataState)
            else ManufacturerDataState(mfg_data_state)
        )
        self.allowed_manufacturers = allowed_manufacturers
        self.is_manufacturer_any = is_manufacturer_any
        self.required_specs = required_specs if required_specs is not None else {}
        self.required_dimensions = required_dimensions if required_dimensions is not None else {}
        self.is_valid = is_valid
        self.validation_errors = validation_errors if validation_errors is not None else []

        # Audit comparison fields
        self.matched_db_row: Optional[Tuple[int, List[str]]] = None
        self.db_part_name: str = ""
        self.db_stock_qty: int = 0
        self.shortage_qty: int = 0
        self.status: str = "NOT_IN_DB"  # "IN_STOCK", "SHORTAGE", "NOT_IN_DB", "INVALID"
        self.link: str = ""

    def update_brand_rule(self, rule_type: str, allowed_brands: Optional[List[str]] = None):
        """
        Updates the in-memory brand rule for this BOM item.
        rule_type: 'ANY', 'EXACT', 'APPROVED_LIST'
        """
        rule_upper = (rule_type or "ANY").upper().strip()
        if rule_upper in ("ANY", "GENERIC"):
            self.manufacturer_rule = ManufacturerRule.ANY
            self.mfg_data_state = ManufacturerDataState.EXPLICIT_ANY
            self.is_manufacturer_any = True
            self.allowed_manufacturers = None
            self.manufacturer = None
            self.approved_manufacturers = []
        elif rule_upper == "APPROVED_LIST" or (allowed_brands and len(allowed_brands) > 1):
            self.manufacturer_rule = ManufacturerRule.APPROVED_LIST
            self.mfg_data_state = ManufacturerDataState.EXPLICIT_BRAND_RULE
            self.is_manufacturer_any = False
            self.allowed_manufacturers = [b.strip() for b in allowed_brands if b.strip()] if allowed_brands else []
            self.approved_manufacturers = list(self.allowed_manufacturers)
            self.manufacturer = ", ".join(self.allowed_manufacturers)
        else:
            self.manufacturer_rule = ManufacturerRule.EXACT
            self.mfg_data_state = ManufacturerDataState.EXPLICIT_BRAND_RULE
            self.is_manufacturer_any = False
            brand_name = allowed_brands[0].strip() if (allowed_brands and len(allowed_brands) > 0) else rule_type.strip()
            self.allowed_manufacturers = [brand_name]
            self.approved_manufacturers = [brand_name]
            self.manufacturer = brand_name

    def _parse_links(self):
        if not self.link or not self.link.strip():
            return None
        try:
            from core.procurement.vendor_links import parse_vendor_links
            from core.procurement.registry import VendorRegistry
            reg = VendorRegistry.get_default_registry()
            parsed = parse_vendor_links(self.link.strip(), registry=reg)

            return parsed
        except Exception:
            return None

    @property
    def link_type(self) -> str:
        """Returns 'LION', 'JAVAN', or vendor-specific string; 'MULTI', 'INVALID', or 'NONE'."""
        if not self.link or not self.link.strip():
            return "NONE"
        parsed = self._parse_links()
        if parsed is None or not parsed.is_valid or parsed.is_empty:
            return "INVALID"
        if parsed.is_empty:
            return "NONE"
        if parsed.count > 1:
            return "MULTI"
        if parsed.count == 1:
            entry = parsed.get_single_entry()
            return entry.vendor_label
        return "INVALID"

    @property
    def is_orderable(self) -> bool:
        """
        Returns True if the line item is eligible for automated procurement.
        In Mode A: Requires a valid registered vendor URL (or multi-vendor links) with an installed and enabled vendor plugin.
        """
        if not self.is_valid or not self.part_number or not self.part_number.strip():
            return False
        if not self.link or not self.link.strip():
            return False
        parsed = self._parse_links()
        if parsed is None or not parsed.is_valid or parsed.is_empty:
            return False
        if parsed.count == 0:
            return False
        try:
            from core.procurement.registry import VendorRegistry
            reg = VendorRegistry.get_default_registry()
            if reg is not None:
                has_active = any(reg.has_vendor(vid) and reg.is_vendor_enabled(vid) for vid in parsed.links.keys())
                return has_active
        except Exception:
            pass
        return parsed.count >= 1

    @property
    def vendor_id(self) -> Optional[str]:
        """Returns the registered vendor identifier if exactly 1 link (or exactly 1 active vendor), or None."""
        if not self.link or not self.link.strip():
            return None
        parsed = self._parse_links()
        if parsed is not None and parsed.is_valid:
            if parsed.count == 1:
                return parsed.get_single_entry().vendor_id
            elif parsed.count > 1:
                try:
                    from core.procurement.registry import VendorRegistry
                    reg = VendorRegistry.get_default_registry()
                    if reg is not None:
                        active_vids = [vid for vid in parsed.links.keys() if reg.has_vendor(vid) and reg.is_vendor_enabled(vid)]
                        if len(active_vids) == 1:
                            return active_vids[0]
                except Exception:
                    pass
        return None

    @property
    def is_multi_vendor(self) -> bool:
        """Returns True if item has multiple active vendor links configured."""
        parsed = self._parse_links()
        if not parsed or not parsed.is_valid or parsed.count <= 1:
            return False
        try:
            from core.procurement.registry import VendorRegistry
            reg = VendorRegistry.get_default_registry()
            if reg is not None:
                active_vids = [vid for vid in parsed.links.keys() if reg.has_vendor(vid) and reg.is_vendor_enabled(vid)]
                return len(active_vids) > 1
        except Exception:
            pass
        return parsed.count > 1

    @property
    def vendor_links(self) -> Dict[str, str]:
        """Returns mapping of vendor_id -> url for all valid links."""
        parsed = self._parse_links()
        if parsed and parsed.is_valid and parsed.count > 0:
            return {vid: e.url for vid, e in parsed.links.items()}
        return {}

    @property
    def link_domain(self) -> str:
        if not self.link or not self.link.strip():
            return ""
        parsed = self._parse_links()
        if parsed and parsed.is_valid:
            if parsed.count > 1:
                return "multi-vendor"
            elif parsed.count == 1:
                try:
                    from urllib.parse import urlparse
                    p = urlparse(parsed.get_single_entry().url)
                    return (p.netloc or parsed.get_single_entry().url).replace("www.", "")
                except Exception:
                    return parsed.get_single_entry().url
        try:
            from urllib.parse import urlparse
            p = urlparse(self.link)
            domain = p.netloc or self.link
            return domain.replace("www.", "")
        except Exception:
            return self.link

    def to_dict(self) -> Dict[str, Any]:
        parsed = self._parse_links()
        effective_link = self.link
        link_error = None
        if parsed is not None:
            if not parsed.is_valid:
                link_error = parsed.error_message
            elif parsed.count == 1:
                effective_link = parsed.get_single_entry().url

        return {
            "row_idx": self.row_idx,
            "designator": self.designator,
            "comment": self.comment,
            "description": self.description,
            "footprint": self.footprint,
            "lib_ref": self.lib_ref,
            "part_number": self.part_number,
            "unit_qty": self.unit_qty,
            "board_multiplier": self.board_multiplier,
            "total_needed": self.total_needed_qty,
            "total_needed_qty": self.total_needed_qty,
            "needed_qty": self.shortage_qty,
            "in_stock": self.db_stock_qty,
            "current_stock": self.db_stock_qty,
            "db_stock_qty": self.db_stock_qty,
            "stock_qty": self.db_stock_qty,
            "shortage": self.shortage_qty,
            "shortage_qty": self.shortage_qty,
            "status": self.status,
            "link": effective_link,
            "raw_link": self.link,
            "link_type": self.link_type,
            "link_domain": self.link_domain,
            "is_orderable": self.is_orderable,
            "is_lion_orderable": self.is_orderable,
            "vendor_id": self.vendor_id,
            "vendor_links": self.vendor_links,
            "is_multi_vendor": self.is_multi_vendor,
            "link_error": link_error,
            "db_part_name": self.db_part_name,
            "supplier_links": getattr(self.matched_db_row, "supplier_links", None) if self.matched_db_row else None,
            "mode": self.mode.value if isinstance(self.mode, ProcurementMode) else str(self.mode),
            "canonical_component_id": self.canonical_component_id,
            "manufacturer": self.manufacturer,
            "approved_manufacturers": self.approved_manufacturers,
            "manufacturer_rule": self.manufacturer_rule.value if self.manufacturer_rule else None,
            "mfg_data_state": self.mfg_data_state.value if isinstance(self.mfg_data_state, ManufacturerDataState) else str(self.mfg_data_state),
            "package": self.package,
            "package_family": self.package_family,
            "pin_count": self.pin_count,
            "value": self.value,
            "tolerance": self.tolerance,
            "voltage_rating": self.voltage_rating,
            "current_rating": self.current_rating,
            "power_rating": self.power_rating,
            "lifecycle_status": self.lifecycle_status,
            "library_version": self.library_version,
            "allowed_manufacturers": self.allowed_manufacturers,
            "is_manufacturer_any": self.is_manufacturer_any,
            "required_specs": self.required_specs,
            "required_dimensions": self.required_dimensions,
            "is_valid": self.is_valid,
            "validation_errors": self.validation_errors,
        }


class BOMAuditResult:
    """Stores the complete analysis outcome of an Altium BOM audit against inventory."""

    def __init__(
        self,
        file_path: str,
        board_multiplier: int,
        mode: ProcurementMode = ProcurementMode.DATABASE_AWARE
    ):
        self.file_path = file_path
        self.file_name = os.path.basename(file_path)
        self.board_multiplier = board_multiplier
        self.mode = mode if isinstance(mode, ProcurementMode) else ProcurementMode(mode)
        self.items: List[BOMItem] = []
        self.total_items_count: int = 0
        self.in_stock_count: int = 0
        self.shortage_count: int = 0
        self.not_in_db_count: int = 0
        self.invalid_count: int = 0
        self.blank_link_count: int = 0

    @property
    def ready_for_assembly(self) -> bool:
        return self.shortage_count == 0 and self.not_in_db_count == 0 and self.invalid_count == 0

    def get_all_shortages(self) -> List[Dict[str, Any]]:
        """Return list of all shortage/orderable items with their procurement classification."""
        shortages = []
        for it in self.items:
            if it.status == "SHORTAGE" and it.shortage_qty > 0 and it.is_valid:
                shortages.append(it.to_dict())
        return shortages

    def get_orderable_shortages(self) -> List[Dict[str, Any]]:
        """Return list of shortage items that are orderable."""
        return [s for s in self.get_all_shortages() if s.get("is_orderable") or s.get("is_lion_orderable")]

    def get_unsupported_shortages(self) -> List[Dict[str, Any]]:
        """Return list of shortage items with external or missing store links / invalid."""
        return [s for s in self.get_all_shortages() if not (s.get("is_orderable") or s.get("is_lion_orderable"))]


class BOMAuditor:
    """
    Parses Altium BOM spreadsheets conforming to DOC-OMNIX-ALTIUM-CONTRACT-001
    and executes deterministic audit under Mode A (Database-Aware Stock Reconciliation).
    """

    @staticmethod
    def audit_bom(
        bom_file_path: str,
        inventory: Optional[Any] = None,
        board_multiplier: int = 1,
        mode: ProcurementMode = ProcurementMode.DATABASE_AWARE,
        deduct_stock: bool = True,
    ) -> Tuple[bool, str, Optional[BOMAuditResult]]:
        """
        Parse Altium BOM Excel file, validate all 20 canonical parameters, and evaluate
        shortages and procurement requirements under Mode A (Database-Aware).
        """
        bom_name = os.path.basename(bom_file_path) if bom_file_path else "Unknown"
        if isinstance(mode, str) and not isinstance(mode, ProcurementMode):
            try:
                mode = ProcurementMode(mode)
            except ValueError:
                raise ValueError(
                    f"Unsupported procurement mode '{mode}'. "
                    "Mode A (DATABASE_AWARE) is the sole supported procurement workflow."
                )
        if mode != ProcurementMode.DATABASE_AWARE:
            raise ValueError(
                f"Unsupported procurement mode '{mode}'. "
                "Mode A (DATABASE_AWARE) is the sole supported procurement workflow."
            )

        tracker = track_operation("BOM", f"BOM audit for '{bom_name}' (Mode: {mode.value}, Multiplier: {board_multiplier})")
        tracker.__enter__()

        if not os.path.exists(bom_file_path):
            msg = f"BOM file not found: {bom_file_path}"
            tracker.fail(msg, level="ERROR")
            return False, msg, None

        try:
            wb = openpyxl.load_workbook(bom_file_path, data_only=True)
            sheet = wb.active or wb[wb.sheetnames[0]]
            raw_rows = list(sheet.iter_rows(values_only=True))
            wb.close()

            if not raw_rows:
                msg = f"BOM file '{bom_name}' is empty."
                tracker.fail(msg, level="ERROR")
                return False, "BOM file is empty.", None

            # Locate columns in BOM header (Row 1)
            raw_headers = [str(h or "").strip() for h in raw_rows[0]]
            header_map = {normalize_string(h): idx for idx, h in enumerate(raw_headers)}

            def find_col(aliases: List[str]) -> Optional[int]:
                for alias in aliases:
                    norm = normalize_string(alias)
                    for h_norm, idx in header_map.items():
                        if norm == h_norm or norm in h_norm:
                            return idx
                return None

            # Extended 20-Parameter Altium Canonical Contract Mapping
            designator_col = find_col(["designator", "محل قطعه", "refdes", "ref"])
            qty_col = find_col(["quantity", "qty", "تعداد"])
            comment_col = find_col(["comment", "کامنت"])
            canonical_id_col = find_col(["canonical_component_id", "component_id", "canonical_id", "comp_id", "id"])
            mpn_col = find_col(["mpn", "part number", "part_number", "part no", "شماره فنی", "mfg part number"])
            mfg_col = find_col(["manufacturer", "mfg", "maker", "brand", "سازنده"])
            approved_mfg_col = find_col([
                "approved_manufacturers", "approved manufacturers", "approved_mfg",
                "approved_brands", "allowed_manufacturers", "سازندگان مجاز"
            ])
            mfg_rule_col = find_col(["manufacturer_rule", "manufacturer rule", "mfg_rule", "brand_rule", "قانون سازنده"])
            package_col = find_col(["package", "case", "پکیج"])
            pkg_family_col = find_col(["package_family", "package family", "pkg_family"])
            footprint_col = find_col(["footprint_id", "footprint", "فوت پرینت", "pcb footprint", "pcb_footprint"])
            pin_count_col = find_col(["pin_count", "pin count", "pins", "pin"])
            desc_col = find_col(["description", "شرح", "desc"])
            value_col = find_col(["value", "مقدار", "val"])
            tolerance_col = find_col(["tolerance", "tol", "تلرانس"])
            voltage_col = find_col(["voltage_rating", "voltage rating", "voltage", "volts", "ولتاژ"])
            current_col = find_col(["current_rating", "current rating", "current", "جریان"])
            power_col = find_col(["power_rating", "power rating", "power", "توان"])
            lifecycle_col = find_col(["lifecycle_status", "lifecycle status", "lifecycle", "status_cad"])
            libref_col = find_col(["library_version", "library version", "lib_version", "libref", "lib ref", "library ref"])

            result = BOMAuditResult(bom_file_path, board_multiplier, mode=mode)

            # Build fast lookup indexes from inventory database
            exact_inv_map: Dict[str, Any] = {}
            is_repo = False
            inv_part_name_col: Optional[int] = None
            inv_part_num_col: Optional[int] = None
            inv_stock_col: Optional[int] = None
            inv_link_col: Optional[int] = None

            if mode == ProcurementMode.DATABASE_AWARE:
                if inventory is None:
                    try:
                        from core.inventory.database import resolve_database_path
                        from core.inventory.repository import SqliteInventoryRepository
                        db_p = resolve_database_path()
                        if db_p and os.path.exists(db_p):
                            inventory = SqliteInventoryRepository(db_p)
                    except Exception:
                        inventory = None

                if inventory is not None:
                    if hasattr(inventory, "get_all_products"):
                        is_repo = True
                        try:
                            for prod in inventory.get_all_products():
                                norm_p = normalize_string(prod.mpn)
                                if norm_p:
                                    exact_inv_map.setdefault(norm_p, []).append(prod)
                        except Exception:
                            pass
                    elif hasattr(inventory, "rows"):
                        inv_part_name_col = inventory.get_column_index(["part name", "نام قطعه", "name"])
                        inv_part_num_col = inventory.get_column_index(["part number", "part_number", "شماره فنی"])
                        inv_stock_col = inventory.get_column_index(["quantity in stock", "stock", "quantity", "موجودی"])
                        inv_link_col = inventory.get_column_index(["link", "url", "لینک"])

                        for orig_row_idx, row_cells in inventory.rows:
                            p_num = (
                                normalize_string(row_cells[inv_part_num_col])
                                if inv_part_num_col is not None and inv_part_num_col < len(row_cells)
                                else ""
                            )
                            p_name = (
                                normalize_string(row_cells[inv_part_name_col])
                                if inv_part_name_col is not None and inv_part_name_col < len(row_cells)
                                else ""
                            )
                            if p_num and p_num not in exact_inv_map:
                                exact_inv_map[p_num] = (orig_row_idx, row_cells)
                            elif p_name and p_name not in exact_inv_map:
                                exact_inv_map[p_name] = (orig_row_idx, row_cells)

            # Parse BOM data rows
            for row_idx, r in enumerate(raw_rows[1:], start=2):
                if not r or all(c is None or str(c).strip() == "" for c in r):
                    continue

                def get_val(col_idx: Optional[int]) -> str:
                    if col_idx is not None and col_idx < len(r) and r[col_idx] is not None:
                        return str(r[col_idx]).strip()
                    return ""

                # Extract 20 parameter values
                designator = get_val(designator_col)
                comment = get_val(comment_col)
                canonical_id = get_val(canonical_id_col)
                raw_mpn = get_val(mpn_col)
                raw_mfg = get_val(mfg_col)
                raw_approved_mfg = get_val(approved_mfg_col)
                raw_mfg_rule = get_val(mfg_rule_col)
                package = get_val(package_col)
                package_family = get_val(pkg_family_col)
                footprint = get_val(footprint_col)
                pin_count_str = get_val(pin_count_col)
                desc = get_val(desc_col)
                value = get_val(value_col)
                tolerance = get_val(tolerance_col)
                voltage = get_val(voltage_col)
                current = get_val(current_col)
                power = get_val(power_col)
                lifecycle = get_val(lifecycle_col)
                lib_ref = get_val(libref_col)
                qty_str = get_val(qty_col)

                validation_errors: List[str] = []
                is_valid = True

                # Strict MPN Validation: Missing MPN -> INVALID_REQUIRED_PARAMETER / Non-orderable
                # Contract Rule: Never fall back to Comment!
                part_number = raw_mpn
                if not part_number:
                    is_valid = False
                    validation_errors.append("INVALID_REQUIRED_PARAMETER: Missing MPN")

                # Strict Quantity Validation
                unit_qty = 0
                if not qty_str:
                    is_valid = False
                    validation_errors.append("INVALID_QUANTITY: Missing quantity")
                else:
                    try:
                        f_qty = float(qty_str)
                        if not f_qty.is_integer() or f_qty <= 0:
                            is_valid = False
                            validation_errors.append(f"INVALID_QUANTITY: Quantity must be a positive integer (got '{qty_str}')")
                        else:
                            unit_qty = int(f_qty)
                    except (ValueError, TypeError):
                        is_valid = False
                        validation_errors.append(f"INVALID_QUANTITY: Invalid quantity string '{qty_str}'")

                # Pin Count Parsing
                pin_count: Optional[int] = None
                if pin_count_str:
                    try:
                        pin_count = int(float(pin_count_str))
                    except (ValueError, TypeError):
                        pin_count = None

                # Manufacturer Rule & Qualification Data State Separation
                # ManufacturerRule: EXACT, APPROVED_LIST, ANY
                # ManufacturerDataState: EXPLICIT_ANY, EXPLICIT_BRAND_RULE, MISSING_MANUFACTURER_DATA
                mfg_rule: Optional[ManufacturerRule] = None
                mfg_data_state: ManufacturerDataState = ManufacturerDataState.MISSING_MANUFACTURER_DATA
                allowed_mfg: Optional[List[str]] = None
                is_mfg_any: bool = False

                # Parse approved manufacturers list
                approved_mfg_list: List[str] = []
                if raw_approved_mfg:
                    approved_mfg_list = [
                        b.strip()
                        for b in re.split(r"[,;|]", raw_approved_mfg)
                        if b.strip()
                    ]

                rule_upper = raw_mfg_rule.upper().strip()
                mfg_upper = raw_mfg.upper().strip()

                if rule_upper in ("ANY", "*", "NONE") or mfg_upper in ("ANY", "*"):
                    mfg_rule = ManufacturerRule.ANY
                    mfg_data_state = ManufacturerDataState.EXPLICIT_ANY
                    is_mfg_any = True
                    allowed_mfg = None
                    raw_mfg = ""
                elif approved_mfg_list or rule_upper in ("APPROVED_LIST", "APPROVED"):
                    mfg_rule = ManufacturerRule.APPROVED_LIST
                    mfg_data_state = ManufacturerDataState.EXPLICIT_BRAND_RULE
                    allowed_mfg = approved_mfg_list if approved_mfg_list else ([raw_mfg] if raw_mfg else None)
                    is_mfg_any = False
                elif raw_mfg:
                    mfg_rule = ManufacturerRule.EXACT
                    mfg_data_state = ManufacturerDataState.EXPLICIT_BRAND_RULE
                    allowed_mfg = [raw_mfg]
                    is_mfg_any = False
                else:
                    mfg_rule = ManufacturerRule.ANY
                    mfg_data_state = ManufacturerDataState.EXPLICIT_ANY
                    allowed_mfg = None
                    is_mfg_any = True
                    raw_mfg = ""

                # Required Specs Bundling
                required_specs: Dict[str, str] = {}
                if value:
                    required_specs["value"] = value
                if tolerance:
                    required_specs["tolerance"] = tolerance
                if voltage:
                    required_specs["voltage_rating"] = voltage
                if current:
                    required_specs["current_rating"] = current
                if power:
                    required_specs["power_rating"] = power

                # Required Dimensions: Populated only when explicit data exists, never fabricated ({})
                required_dimensions: Dict[str, str] = {}

                bom_item = BOMItem(
                    row_idx=row_idx,
                    designator=designator,
                    comment=comment,
                    description=desc,
                    footprint=footprint,
                    lib_ref=lib_ref,
                    part_number=part_number,
                    unit_qty=unit_qty,
                    board_multiplier=board_multiplier,
                    canonical_component_id=canonical_id,
                    manufacturer=raw_mfg,
                    approved_manufacturers=approved_mfg_list,
                    manufacturer_rule=mfg_rule,
                    package=package,
                    package_family=package_family,
                    pin_count=pin_count,
                    value=value,
                    tolerance=tolerance,
                    voltage_rating=voltage,
                    current_rating=current,
                    power_rating=power,
                    lifecycle_status=lifecycle,
                    library_version=lib_ref,
                    mode=mode,
                    mfg_data_state=mfg_data_state,
                    allowed_manufacturers=allowed_mfg,
                    is_manufacturer_any=is_mfg_any,
                    required_specs=required_specs,
                    required_dimensions=required_dimensions,
                    is_valid=is_valid,
                    validation_errors=validation_errors,
                )

                if not is_valid:
                    bom_item.status = "INVALID"
                    bom_item.db_stock_qty = 0
                    bom_item.shortage_qty = bom_item.total_needed_qty
                    result.invalid_count += 1
                    result.items.append(bom_item)
                    continue

                # Mode A: Database-Aware Shortage Reconciliation
                norm_part_num = normalize_string(part_number)

                if is_repo:
                    matched_prod = None
                    if norm_part_num and norm_part_num in exact_inv_map:
                        prods = exact_inv_map[norm_part_num]
                        if len(prods) == 1:
                            matched_prod = prods[0]
                        else:
                            norm_bom_mfg = normalize_string(raw_mfg)
                            for p in prods:
                                if norm_bom_mfg and normalize_string(p.manufacturer) == norm_bom_mfg:
                                    matched_prod = p
                                    break
                            if not matched_prod:
                                matched_prod = prods[0]

                    if matched_prod:
                        bom_item.matched_db_row = matched_prod
                        bom_item.db_part_name = matched_prod.mpn
                        bom_item.db_stock_qty = int(matched_prod.quantity_in_stock or 0)
                        bom_item.link = matched_prod.purchase_link or ""

                        if deduct_stock:
                            if bom_item.db_stock_qty >= bom_item.total_needed_qty:
                                bom_item.status = "IN_STOCK"
                                bom_item.shortage_qty = 0
                                result.in_stock_count += 1
                            else:
                                bom_item.status = "SHORTAGE"
                                bom_item.shortage_qty = bom_item.total_needed_qty - bom_item.db_stock_qty
                                result.shortage_count += 1
                                if not bom_item.link or not bom_item.link.strip():
                                    result.blank_link_count += 1
                        else:
                            bom_item.status = "SHORTAGE"
                            bom_item.shortage_qty = bom_item.total_needed_qty
                            result.shortage_count += 1
                            if not bom_item.link or not bom_item.link.strip():
                                result.blank_link_count += 1
                    else:
                        bom_item.status = "NOT_IN_DB"
                        bom_item.db_stock_qty = 0
                        bom_item.shortage_qty = bom_item.total_needed_qty
                        bom_item.link = ""
                        result.not_in_db_count += 1

                else:
                    # Legacy ExcelDataLoader path: strictly match MPN without Comment/Description fallbacks
                    match_found = None
                    if norm_part_num and norm_part_num in exact_inv_map:
                        match_found = exact_inv_map[norm_part_num]

                    if match_found:
                        orig_idx, inv_cells = match_found
                        bom_item.matched_db_row = match_found
                        bom_item.db_part_name = (
                            inv_cells[inv_part_name_col]
                            if inv_part_name_col is not None and inv_part_name_col < len(inv_cells)
                            else ""
                        )

                        stock_str = (
                            inv_cells[inv_stock_col]
                            if inv_stock_col is not None and inv_stock_col < len(inv_cells)
                            else "0"
                        )
                        try:
                            bom_item.db_stock_qty = int(float(stock_str)) if stock_str else 0
                        except ValueError:
                            bom_item.db_stock_qty = 0

                        bom_item.link = (
                            inv_cells[inv_link_col]
                            if inv_link_col is not None and inv_link_col < len(inv_cells)
                            else ""
                        )

                        if deduct_stock:
                            if bom_item.db_stock_qty >= bom_item.total_needed_qty:
                                bom_item.status = "IN_STOCK"
                                bom_item.shortage_qty = 0
                                result.in_stock_count += 1
                            else:
                                bom_item.status = "SHORTAGE"
                                bom_item.shortage_qty = bom_item.total_needed_qty - bom_item.db_stock_qty
                                result.shortage_count += 1
                        else:
                            bom_item.status = "SHORTAGE"
                            bom_item.shortage_qty = bom_item.total_needed_qty
                            result.shortage_count += 1
                    else:
                        bom_item.status = "NOT_IN_DB"
                        bom_item.db_stock_qty = 0
                        bom_item.shortage_qty = bom_item.total_needed_qty
                        bom_item.link = ""
                        result.not_in_db_count += 1

                result.items.append(bom_item)

            result.total_items_count = len(result.items)
            tracker.complete(
                f"Completed BOM audit for '{bom_name}': {result.total_items_count} items evaluated ({result.shortage_count} shortages, {result.in_stock_count} in stock)"
            )
            return True, f"Audited {result.total_items_count} components from BOM ({mode.value}).", result

        except Exception as e:
            err_msg = f"Failed to audit BOM '{bom_name}': {str(e)}"
            tracker.fail(err_msg, level="ERROR")
            return False, f"Failed to audit BOM: {str(e)}", None
